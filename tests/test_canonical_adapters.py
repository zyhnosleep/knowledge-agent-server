from __future__ import annotations

import shutil
from base64 import b64decode
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from lxml import html as lxml_html

from app.services.canonical_artifacts import CanonicalArtifactStore
from app.services import canonical_adapters
from app.services.canonical_adapters import (
    DocxCanonicalAdapter,
    PDFCanonicalAdapter,
    parse_canonical_document,
)
from app.services.parser import DocumentParseError, parse_document


FIXTURE_DIR = Path(__file__).parent / "fixtures" / "canonical"


@pytest.fixture(autouse=True)
def _isolate_adapter_asset_cache(monkeypatch, tmp_path: Path) -> None:
    class TestSettings:
        cache_dir = tmp_path / "adapter-cache"

    monkeypatch.setattr(canonical_adapters, "get_settings", lambda: TestSettings())


def _append_omml(paragraph, text: str = "x + y = 1") -> None:
    formula = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text_node = OxmlElement("m:t")
    text_node.set(qn("xml:space"), "preserve")
    text_node.text = text
    run.append(text_node)
    formula.append(run)
    paragraph._p.append(formula)


def _write_png(path: Path) -> None:
    path.write_bytes(
        b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/"
            "x8AAusB9Wl2nKsAAAAASUVORK5CYII="
        )
    )


def _write_structured_docx(path: Path, image_path: Path) -> None:
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("before ")
    _append_omml(paragraph)
    paragraph.add_run(" after ")
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(0.1))
    shape._inline.docPr.set("descr", "Inline chart")
    paragraph.add_run(" tail")

    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "Scores"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "Recall"
    table.cell(1, 1).text = "0.91"
    media_cell = table.cell(1, 2)
    media_cell.text = "Evidence"
    media_shape = media_cell.paragraphs[0].add_run().add_picture(
        str(image_path), width=Inches(0.1)
    )
    media_shape._inline.docPr.set("descr", "Cell chart")
    _append_omml(media_cell.paragraphs[0], "z = 2")
    table.cell(2, 1).merge(table.cell(2, 2)).text = "Aggregate"
    document.save(path)


def _write_wrapped_nested_docx(path: Path) -> None:
    document = Document()
    wrapped_paragraph = document.add_paragraph("Wrapped paragraph.")
    wrapped_table = document.add_table(rows=1, cols=1)
    wrapped_table.cell(0, 0).text = "Wrapped table"
    body = document.element.body
    body.remove(wrapped_paragraph._p)
    body.remove(wrapped_table._tbl)
    content_control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    content.append(wrapped_paragraph._p)
    content.append(wrapped_table._tbl)
    content_control.append(content)
    body.insert(0, content_control)

    outer = document.add_table(rows=1, cols=1)
    cell = outer.cell(0, 0)
    cell.paragraphs[0].add_run("before")
    _append_omml(cell.paragraphs[0], "before = 1")
    nested = cell.add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "Nested table"
    after = cell.add_paragraph()
    _append_omml(after, "after = 2")
    document.save(path)


def _assert_common_contract(document) -> None:
    assert document.blocks
    assert document.source_path
    assert document.source_media_type
    assert document.parser_source
    assert document.parser_metadata
    assert [block.reading_order for block in document.blocks] == list(
        range(len(document.blocks))
    )
    block_ids = [block.block_id for block in document.blocks]
    assert all(block_ids)
    assert len(block_ids) == len(set(block_ids))
    assert all(block.parser_source for block in document.blocks)

    table_ids = {table.table_id for table in document.tables}
    figure_ids = {figure.figure_id for figure in document.figures}
    formula_ids = {formula.formula_id for formula in document.formulas}
    assert all(block.table_id in table_ids for block in document.blocks if block.table_id)
    assert all(block.figure_id in figure_ids for block in document.blocks if block.figure_id)
    assert all(block.formula_id in formula_ids for block in document.blocks if block.formula_id)


@pytest.mark.parametrize("name", ["sample.md", "sample.html", "sample.docx"])
def test_existing_formats_produce_canonical_blocks(name: str) -> None:
    document = parse_canonical_document(FIXTURE_DIR / name)
    _assert_common_contract(document)


def test_block_ids_are_stable_across_repeated_parse() -> None:
    path = FIXTURE_DIR / "sample.md"
    first = parse_canonical_document(path)
    second = parse_canonical_document(path)
    assert [block.block_id for block in first.blocks] == [block.block_id for block in second.blocks]


def test_markdown_preserves_structure_and_exact_source_ranges() -> None:
    path = FIXTURE_DIR / "sample.md"
    source = path.read_text(encoding="utf-8")
    document = parse_canonical_document(path)

    intro = next(block for block in document.blocks if block.text.startswith("This paragraph"))
    assert intro.section_path == ["Canonical Adapter Study", "Introduction"]
    assert source[intro.source_spans[0].char_start : intro.source_spans[0].char_end] == intro.text

    headings = [block for block in document.blocks if block.block_type == "heading"]
    assert headings and all(not block.retrievable for block in headings)
    assert document.outline[0].title == "Canonical Adapter Study"
    assert document.outline[0].children[0].title == "Introduction"

    code = next(block for block in document.blocks if block.metadata.get("kind") == "code")
    code_source = source[code.source_spans[0].char_start : code.source_spans[0].char_end]
    assert code.metadata["language"] == "python"
    assert code.metadata["fence"] == "```"
    assert code_source.startswith("```python") and code_source.endswith("```")
    assert "return value" in code.text

    table_block = next(block for block in document.blocks if block.block_type == "table")
    table = next(table for table in document.tables if table.table_id == table_block.table_id)
    assert table.headers == ["Metric", "Score"]
    assert table.rows == [["Recall", "0.91"], ["Precision", "0.88"]]
    assert table.source_markdown == source[
        table_block.source_spans[0].char_start : table_block.source_spans[0].char_end
    ]
    assert len(table.cells) == 6

    formula_block = next(block for block in document.blocks if block.block_type == "formula")
    formula = next(item for item in document.formulas if item.formula_id == formula_block.formula_id)
    formula_source = source[
        formula_block.source_spans[0].char_start : formula_block.source_spans[0].char_end
    ]
    assert formula.latex == r"F_1 = 2 \frac{PR}{P + R}"
    assert formula_source.startswith("$$") and formula_source.endswith("$$")

    figure_block = next(block for block in document.blocks if block.block_type == "figure")
    figure = next(item for item in document.figures if item.figure_id == figure_block.figure_id)
    assert figure.caption == "Evaluation curve"
    assert figure.metadata["target"] == "figures/evaluation.png"
    assert source[
        figure_block.source_spans[0].char_start : figure_block.source_spans[0].char_end
    ].startswith("![Evaluation curve]")


def test_markdown_references_stop_at_appendix_and_appendix_is_retrievable() -> None:
    document = parse_canonical_document(FIXTURE_DIR / "sample.md")
    reference = next(block for block in document.blocks if block.text.startswith("Doe, J."))
    assert reference.block_type == "reference"
    assert reference.retrievable is False

    appendix = next(block for block in document.blocks if block.block_type == "appendix")
    assert appendix.retrievable is True
    assert appendix.section_path[-1] == "Appendix"


def test_reference_state_applies_to_every_structure_until_peer_heading(tmp_path: Path) -> None:
    path = tmp_path / "references.md"
    path.write_text(
        """# Study

## References

Reference prose.

```text
reference code
```

| Source | Year |
| --- | --- |
| Paper | 2026 |

$$
r = 1
$$

![Reference figure](ref.png)

### Appendix

Nested reference prose.

## Appendix

Appendix prose.
""",
        encoding="utf-8",
    )
    document = parse_canonical_document(path)
    appendix_heading = [
        block
        for block in document.blocks
        if block.block_type == "heading" and block.text == "Appendix"
    ][-1]
    references = [
        block
        for block in document.blocks
        if block.reading_order < appendix_heading.reading_order and block.block_type != "heading"
    ]
    assert [block.block_type for block in references] == [
        "reference",
        "reference",
        "table",
        "formula",
        "figure",
        "reference",
    ]
    assert all(not block.retrievable for block in references)
    assert references[1].metadata["kind"] == "code"
    nested_reference = next(
        block for block in document.blocks if block.text == "Nested reference prose."
    )
    assert nested_reference.section_path[-1] == "Appendix"
    appendix = next(block for block in document.blocks if block.text == "Appendix prose.")
    assert appendix.retrievable is True

    html_path = tmp_path / "reference-caption.html"
    html_path.write_text(
        """<h1>Study</h1><h2>References</h2>
<table><caption>Reference caption</caption><tr><td>Source</td></tr></table>
<figure><img src="ref.png"><figcaption>Figure caption</figcaption></figure>
<h2>Appendix</h2><p>Recovered</p>""",
        encoding="utf-8",
    )
    html_document = parse_canonical_document(html_path)
    reference_structures = [
        block
        for block in html_document.blocks
        if block.block_type in {"caption", "table", "figure"}
    ]
    assert reference_structures and all(
        not block.retrievable for block in reference_structures
    )
    assert next(block for block in html_document.blocks if block.text == "Recovered").retrievable


def test_markdown_setext_inline_image_and_fence_payload_are_exact(tmp_path: Path) -> None:
    path = tmp_path / "structures.md"
    source = (
        "Setext Title\r\n"
        "============\r\n"
        "\r\n"
        "before ![Inline alt](<images/chart one.png> \"Chart title\") after\r\n"
        "\r\n"
        "```text\r\n"
        "value\r\n"
        "\r\n"
        "\r\n"
        "```\r\n"
        "\r\n"
        "~~~raw\r\n"
        "unclosed\r\n"
    )
    path.write_bytes(source.encode("utf-8"))
    document = parse_canonical_document(path)
    assert document.title == "Setext Title"
    assert document.blocks[0].metadata["heading_level"] == 1
    inline = [block for block in document.blocks if 1 <= block.reading_order <= 3]
    assert [block.block_type for block in inline] == ["narrative", "figure", "narrative"]
    assert [block.text.strip() for block in inline] == ["before", "Inline alt", "after"]
    figure = document.figures[0]
    assert figure.asset_path == "images/chart one.png"
    assert figure.metadata["title"] == "Chart title"
    for block in inline:
        span = block.source_spans[0]
        assert source[span.char_start : span.char_end] == block.metadata.get(
            "source_markdown", block.text
        )

    closed, unclosed = [
        block for block in document.blocks if block.metadata.get("kind") == "code"
    ]
    assert closed.text == "value\r\n\r\n\r\n"
    assert unclosed.text == "unclosed\r\n"
    assert unclosed.metadata["closed"] is False
    unclosed_span = unclosed.source_spans[0]
    assert source[unclosed_span.char_start : unclosed_span.char_end] == (
        "~~~raw\r\nunclosed\r\n"
    )


def test_markdown_commonmark_image_scanner_atx_and_fence_info(tmp_path: Path) -> None:
    path = tmp_path / "images.md"
    source = """# title#

`![code](skip-code.png)` and \\![escaped](skip-escape.png).

`code `` ![Still code](skip-long-run.png) `

before ![Balanced](plots/a_(b).png "Direct") after

![Nested [alt]](nested.png)

![Reference][chart]

![Collapsed][]

[chart]: <images/reference chart.png> "Reference title"
[Collapsed]: collapsed.png

```python linenos
print("ok")
```
    """
    path.write_text(source, encoding="utf-8")
    with path.open("r", encoding="utf-8", newline="") as source_file:
        source = source_file.read()
    document = parse_canonical_document(path)
    assert document.blocks[0].text == "title#"
    assert [figure.asset_path for figure in document.figures] == [
        "plots/a_(b).png",
        "nested.png",
        "images/reference chart.png",
        "collapsed.png",
    ]
    assert [figure.caption for figure in document.figures] == [
        "Balanced",
        "Nested [alt]",
        "Reference",
        "Collapsed",
    ]
    assert all("skip-" not in figure.asset_path for figure in document.figures)
    code = next(block for block in document.blocks if block.metadata.get("kind") == "code")
    assert code.metadata["language"] == "python"
    assert code.metadata["info"] == "python linenos"
    for block in [item for item in document.blocks if item.block_type == "figure"]:
        span = block.source_spans[0]
        assert source[span.char_start : span.char_end] == block.metadata["source_markdown"]


@pytest.mark.parametrize("suffix", [".md", ".txt"])
def test_bom_crlf_non_ascii_source_spans_use_original_character_stream(
    tmp_path: Path, suffix: str
) -> None:
    path = tmp_path / f"bom{suffix}"
    body = "# 标题\r\n\r\n中文段落。" if suffix == ".md" else "中文第一段。\r\n\r\n第二段。"
    path.write_bytes(b"\xef\xbb\xbf" + body.encode("utf-8"))
    with path.open("r", encoding="utf-8", newline="") as source_file:
        source = source_file.read()
    document = parse_canonical_document(path)
    assert document.blocks[0].source_spans[0].char_start == 1
    for block in document.blocks:
        span = block.source_spans[0]
        excerpt = source[span.char_start : span.char_end]
        if block.block_type == "heading":
            assert excerpt.startswith("# ") and block.text in excerpt
        else:
            assert excerpt == block.text


def test_html_preserves_dom_order_structures_and_locators() -> None:
    document = parse_canonical_document(FIXTURE_DIR / "sample.html")
    _assert_common_contract(document)
    combined = "\n".join(block.text for block in document.blocks)
    assert "script-content" not in combined
    assert "Navigation content" not in combined
    assert "noscript-content" not in combined
    assert "template-content" not in combined
    assert "Rendered loss" not in combined
    assert document.title == "Canonical HTML Study"

    summary = next(block for block in document.blocks if block.text.startswith("The adapter"))
    assert summary.section_path == ["Canonical HTML Study", "Methods"]
    assert summary.source_spans[0].element_id == "method-summary"
    assert summary.source_spans[0].css_selector == "#method-summary"
    assert summary.source_spans[0].xpath

    code = next(block for block in document.blocks if block.metadata.get("kind") == "code")
    assert code.text == "result = parse(source)"
    assert code.source_spans[0].css_selector == "#example-code"

    table = document.tables[0]
    assert table.headers == ["Metric", "Score", "", ""]
    assert table.rows == [["", "Recall", "0.91", "high"]]
    assert any(cell.rowspan == 2 for cell in table.cells)
    assert any(cell.colspan == 2 for cell in table.cells)
    assert table.source_html.startswith("<table")
    assert table.source_spans[0].css_selector == "#results-table"

    figure = document.figures[0]
    assert figure.caption == "Figure 1. Evaluation curve"
    assert figure.metadata["src"] == "figures/evaluation.png"
    assert figure.metadata["alt"] == "Evaluation curve"
    assert figure.source_spans[0].css_selector == "#evaluation-figure"

    formula = document.formulas[0]
    assert formula.latex == r"L = -\log p(y|x)"
    assert formula.source_spans[0].css_selector == "#loss-equation"


def test_html_inline_content_is_split_in_exact_dom_order_without_parent_duplication(
    tmp_path: Path,
) -> None:
    path = tmp_path / "inline.html"
    path.write_text(
        '<html><body><h1>Inline</h1>\n<p id="mixed">before '
        '<code>code()</code> middle <span id="inline-math" class="math" '
        'data-latex="x=1">rendered</span> after <img id="inline-image" '
        'src="chart.png" alt="Chart"> tail</p>\n</body></html>',
        encoding="utf-8",
    )
    document = parse_canonical_document(path)
    blocks = document.blocks[1:]
    assert [block.block_type for block in blocks] == [
        "narrative",
        "narrative",
        "narrative",
        "formula",
        "narrative",
        "figure",
        "narrative",
    ]
    assert [block.text.strip() for block in blocks] == [
        "before",
        "code()",
        "middle",
        "x=1",
        "after",
        "Chart",
        "tail",
    ]
    assert blocks[1].metadata["kind"] == "code"
    assert "rendered" not in "\n".join(block.text for block in blocks)
    assert blocks[3].source_spans[0].css_selector == "#inline-math"
    assert blocks[5].source_spans[0].css_selector == "#inline-image"


def test_html_captions_nested_tables_and_spans_have_structural_coordinates(
    tmp_path: Path,
) -> None:
    path = tmp_path / "tables.html"
    path.write_text(
        """<html><body><h1>Tables</h1>
<table id="outer"><caption id="outer-caption">Outer results</caption>
<tr><th>A</th><th colspan="2">B</th></tr>
<tr><td rowspan="2">R</td><td>1</td><td>outer<table id="inner">
<caption id="inner-caption">Inner results</caption><tr><th>I</th></tr>
<tr><td>2</td></tr></table></td></tr>
<tr><td colspan="2">Aggregate</td></tr></table>
<figure id="figure"><img src="chart.png" alt="Chart">
<figcaption id="figure-caption">Figure caption</figcaption></figure>
<table id="malformed"><tr><td rowspan="bad" colspan="nope">safe</td></tr></table>
</body></html>""",
        encoding="utf-8",
    )
    document = parse_canonical_document(path)
    outer, inner, malformed = document.tables
    assert outer.headers == ["A", "B", ""]
    assert outer.rows == [["R", "1", "outer"], ["", "Aggregate", ""]]
    assert {
        (cell.text, cell.row_index, cell.column_index, cell.rowspan, cell.colspan)
        for cell in outer.cells
    } == {
        ("A", 0, 0, 1, 1),
        ("B", 0, 1, 1, 2),
        ("R", 1, 0, 2, 1),
        ("1", 1, 1, 1, 1),
        ("outer", 1, 2, 1, 1),
        ("Aggregate", 2, 1, 1, 2),
    }
    assert inner.headers == ["I"] and inner.rows == [["2"]]
    assert malformed.cells[0].rowspan == 1 and malformed.cells[0].colspan == 1
    assert any("rowspan" in warning or "colspan" in warning for warning in document.warnings)

    captions = [block for block in document.blocks if block.block_type == "caption"]
    outer_caption = next(block for block in captions if block.text == "Outer results")
    figure_caption = next(block for block in captions if block.text == "Figure caption")
    assert outer_caption.table_id == outer.table_id
    assert outer_caption.source_spans[0].css_selector == "#outer-caption"
    assert figure_caption.figure_id == document.figures[0].figure_id
    assert figure_caption.source_spans[0].css_selector == "#figure-caption"
    assert len([block for block in document.blocks if block.table_id == outer.table_id]) == 2


def test_html_colspan_moves_past_every_active_rowspan_column(tmp_path: Path) -> None:
    path = tmp_path / "overlap.html"
    path.write_text(
        """<table><tr><td>A</td><td rowspan="2">Held</td></tr>
<tr><td colspan="2">Wide</td></tr></table>""",
        encoding="utf-8",
    )
    table = parse_canonical_document(path).tables[0]
    wide = next(cell for cell in table.cells if cell.text == "Wide")
    assert (wide.row_index, wide.column_index, wide.colspan) == (1, 2, 2)
    assert table.rows[1] == ["", "", "Wide", ""]


def test_html_table_and_figure_children_preserve_order_and_parent_links(
    tmp_path: Path,
) -> None:
    path = tmp_path / "nested-structures.html"
    path.write_text(
        """<html><body><table id="parent-table"><tr><td>
<code id="cell-code">cell()</code><span id="cell-math" class="math" data-latex="c=1">render</span>
<img id="cell-image" src="cell.png" alt="Cell image">
<table id="cell-table"><tr><td>nested</td></tr></table>
</td></tr></table>
<figure id="rich-figure"><img id="main-image" src="main.png" alt="Main alt">
<code id="figure-code">figure()</code><span id="figure-math" class="math" data-latex="f=1">render</span>
<img id="extra-image" src="extra.png" alt="Extra alt">
<table id="figure-table"><tr><td>inside</td></tr></table>
<figcaption id="rich-caption">Rich caption</figcaption></figure></body></html>""",
        encoding="utf-8",
    )
    document = parse_canonical_document(path)
    parent_table = next(table for table in document.tables if table.source_spans[0].element_id == "parent-table")
    parent_children = [
        block
        for block in document.blocks
        if block.metadata.get("parent_table_id") == parent_table.table_id
    ]
    assert [block.block_type for block in parent_children] == [
        "narrative",
        "formula",
        "figure",
        "table",
    ]
    assert parent_children[0].metadata["kind"] == "code"
    assert all(block.source_spans[0].metadata["parent_table_id"] == parent_table.table_id for block in parent_children)

    rich_figure = next(figure for figure in document.figures if figure.source_spans[0].element_id == "rich-figure")
    figure_block = next(block for block in document.blocks if block.figure_id == rich_figure.figure_id)
    assert figure_block.text == "Main alt"
    assert "Rich caption" not in figure_block.text
    figure_children = [
        block
        for block in document.blocks
        if block.metadata.get("parent_figure_id") == rich_figure.figure_id
    ]
    assert [block.block_type for block in figure_children] == [
        "narrative",
        "formula",
        "figure",
        "table",
    ]
    assert len([figure for figure in document.figures if figure.metadata.get("src") == "main.png"]) == 1
    caption = next(block for block in document.blocks if block.text == "Rich caption")
    assert caption.block_type == "caption" and caption.figure_id == rich_figure.figure_id


def test_html_rowspan_zero_stops_at_direct_row_group_boundary(tmp_path: Path) -> None:
    path = tmp_path / "rowspan-zero.html"
    path.write_text(
        """<table><tbody>
<tr><td rowspan="0">Body</td><td>1</td></tr><tr><td>2</td></tr><tr><td>3</td></tr>
</tbody><tfoot><tr><td>Foot</td><td>4</td></tr></tfoot></table>""",
        encoding="utf-8",
    )
    table = parse_canonical_document(path).tables[0]
    body = next(cell for cell in table.cells if cell.text == "Body")
    foot = next(cell for cell in table.cells if cell.text == "Foot")
    assert body.rowspan == 3
    assert (foot.row_index, foot.column_index) == (3, 0)
    assert table.rows == [
        ["Body", "1"],
        ["", "2"],
        ["", "3"],
        ["Foot", "4"],
    ]


def test_html_text_segment_locator_points_to_real_container(tmp_path: Path) -> None:
    path = tmp_path / "locator.html"
    source = '<html><body><p id="locator">before <em>emphasis</em> after</p></body></html>'
    path.write_text(source, encoding="utf-8")
    block = parse_canonical_document(path).blocks[0]
    span = block.source_spans[0]
    selected = lxml_html.fromstring(source).getroottree().xpath(span.xpath)
    assert selected and selected[0].tag == "p"
    assert span.css_selector == "#locator"
    assert span.metadata["segment_index"] == 0


def test_docx_preserves_body_order_locators_image_and_omml() -> None:
    document = parse_canonical_document(FIXTURE_DIR / "sample.docx")
    _assert_common_contract(document)
    block_types = [block.block_type for block in document.blocks]
    paragraph_index = next(
        index for index, block in enumerate(document.blocks) if block.text == "Paragraph before table."
    )
    table_index = block_types.index("table")
    after_index = next(
        index for index, block in enumerate(document.blocks) if block.text == "Paragraph after table."
    )
    assert paragraph_index < table_index < after_index

    heading = next(block for block in document.blocks if block.block_type == "heading")
    assert heading.text == "DOCX Methods"
    assert heading.source_spans[0].paragraph_id
    paragraph = next(block for block in document.blocks if block.text == "Paragraph before table.")
    assert paragraph.section_path == ["DOCX Methods"]
    assert paragraph.source_spans[0].paragraph_id

    table = document.tables[0]
    assert table.headers == ["Metric", "Score"]
    assert table.rows == [["Recall", "0.91"]]
    assert all(cell.source_spans[0].table_id == table.table_id for cell in table.cells)
    assert {(cell.source_spans[0].row_index, cell.source_spans[0].column_index) for cell in table.cells} == {
        (0, 0), (0, 1), (1, 0), (1, 1)
    }

    figure = document.figures[0]
    assert figure.source_spans[0].image_relationship_id
    assert figure.asset_path
    assert document.assets and document.assets[0].media_type == "image/png"

    formula = document.formulas[0]
    assert "x + y = 1" in formula.latex
    assert formula.metadata["source_format"] == "omml"
    assert "oMath" in formula.metadata["omml"]


def test_docx_inline_events_and_merged_table_cells_preserve_structure(tmp_path: Path) -> None:
    image_path = tmp_path / "pixel.png"
    docx_path = tmp_path / "structured.docx"
    _write_png(image_path)
    _write_structured_docx(docx_path, image_path)
    document = DocxCanonicalAdapter(asset_cache_root=tmp_path / "asset-cache").parse(docx_path)

    table_order = next(block.reading_order for block in document.blocks if block.block_type == "table")
    inline_blocks = [block for block in document.blocks if block.reading_order < table_order]
    assert [block.block_type for block in inline_blocks] == [
        "narrative",
        "formula",
        "narrative",
        "figure",
        "narrative",
    ]
    assert [block.text.strip() for block in inline_blocks] == [
        "before",
        "x + y = 1",
        "after",
        "Inline chart",
        "tail",
    ]
    assert "before" not in document.figures[0].caption

    table = document.tables[0]
    coordinates = {
        (cell.text, cell.row_index, cell.column_index, cell.rowspan, cell.colspan)
        for cell in table.cells
    }
    assert ("Scores", 0, 1, 1, 2) in coordinates
    assert ("Recall", 1, 0, 2, 1) in coordinates
    assert len([cell for cell in table.cells if cell.text == "Scores"]) == 1
    assert len([cell for cell in table.cells if cell.text == "Recall"]) == 1
    assert table.headers == ["Metric", "Scores", ""]
    assert table.rows[0][:2] == ["Recall", "0.91"]
    assert table.rows[1] == ["", "Aggregate", ""]

    cell_figures = [
        block
        for block in document.blocks
        if block.block_type == "figure" and block.source_spans[0].table_id == table.table_id
    ]
    cell_formulas = [
        block
        for block in document.blocks
        if block.block_type == "formula" and block.source_spans[0].table_id == table.table_id
    ]
    assert len(cell_figures) == 1 and len(cell_formulas) == 1
    assert (cell_figures[0].source_spans[0].row_index, cell_figures[0].source_spans[0].column_index) == (1, 2)
    assert cell_formulas[0].text == "z = 2"


def test_docx_assets_can_be_promoted_and_loaded_without_touching_source_directory(
    tmp_path: Path,
) -> None:
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    docx_path = source_dir / "sample.docx"
    shutil.copy2(FIXTURE_DIR / "sample.docx", docx_path)
    before = {item.name for item in source_dir.iterdir()}

    document = DocxCanonicalAdapter(asset_cache_root=tmp_path / "asset-cache").parse(docx_path)
    assert {item.name for item in source_dir.iterdir()} == before
    asset = document.assets[0]
    assert asset.path.startswith("assets/")
    assert Path(asset.source_path).is_file()
    assert Path(asset.source_path).parent != source_dir

    store = CanonicalArtifactStore(tmp_path / "artifacts")
    staging = store.write_staging(document.document_id, document.parse_version, document)
    assert (staging / asset.path).is_file()
    final = store.promote(document.document_id, document.parse_version)
    loaded = store.load(document.document_id, document.parse_version)
    assert (final / loaded.assets[0].path).is_file()
    assert loaded.assets[0].sha256 == asset.sha256


def test_docx_wrappers_and_cell_nested_tables_follow_recursive_xml_order(
    tmp_path: Path,
) -> None:
    path = tmp_path / "wrapped.docx"
    _write_wrapped_nested_docx(path)
    document = DocxCanonicalAdapter(asset_cache_root=tmp_path / "asset-cache").parse(path)
    wrapped_text = next(block for block in document.blocks if block.text == "Wrapped paragraph.")
    wrapped_table = next(table for table in document.tables if "Wrapped table" in table.headers)
    wrapped_table_block = next(block for block in document.blocks if block.table_id == wrapped_table.table_id)
    assert wrapped_text.reading_order < wrapped_table_block.reading_order

    outer = next(table for table in document.tables if table.metadata.get("locator") == "body/1")
    nested = next(table for table in document.tables if table.metadata.get("parent_table_id") == outer.table_id)
    outer_block = next(block for block in document.blocks if block.table_id == outer.table_id)
    nested_block = next(block for block in document.blocks if block.table_id == nested.table_id)
    before = next(block for block in document.blocks if block.text == "before = 1")
    after = next(block for block in document.blocks if block.text == "after = 2")
    assert outer_block.reading_order < before.reading_order < nested_block.reading_order < after.reading_order
    assert nested.metadata["parent_row_index"] == 0
    assert nested.metadata["parent_column_index"] == 0
    assert nested.metadata["locator"].startswith("body/1/cell-0-0/")
    assert nested_block.metadata["parent_table_id"] == outer.table_id
    assert nested_block.source_spans[0].metadata["parent_table_id"] == outer.table_id


@pytest.mark.parametrize("name", ["sample.md", "sample.html", "sample.docx"])
def test_canonical_resource_ids_are_stable_after_source_relocation(
    tmp_path: Path, name: str
) -> None:
    first_dir = tmp_path / "first"
    second_dir = tmp_path / "second"
    first_dir.mkdir()
    second_dir.mkdir()
    first_path = first_dir / name
    second_path = second_dir / name
    shutil.copy2(FIXTURE_DIR / name, first_path)
    shutil.copy2(FIXTURE_DIR / name, second_path)
    if name.endswith(".docx"):
        adapter = DocxCanonicalAdapter(asset_cache_root=tmp_path / "asset-cache")
        first = adapter.parse(first_path)
        second = adapter.parse(second_path)
    else:
        first = parse_canonical_document(first_path)
        second = parse_canonical_document(second_path)

    def ids(document) -> tuple:
        return (
            document.document_id,
            [block.block_id for block in document.blocks],
            [table.table_id for table in document.tables],
            [figure.figure_id for figure in document.figures],
            [formula.formula_id for formula in document.formulas],
            [asset.asset_id for asset in document.assets],
        )

    assert ids(first) == ids(second)


def test_text_and_unknown_suffix_preserve_exact_ranges_and_empty_file_warning(tmp_path: Path) -> None:
    source = "First paragraph.\ncontinues here.\n\nSecond paragraph stays whole.\n"
    text_path = tmp_path / "notes.txt"
    text_path.write_text(source, encoding="utf-8")
    with text_path.open("r", encoding="utf-8", newline="") as source_file:
        source = source_file.read()
    text_document = parse_canonical_document(text_path)
    assert [block.text.replace("\r\n", "\n") for block in text_document.blocks] == [
        "First paragraph.\ncontinues here.",
        "Second paragraph stays whole.",
    ]
    for block in text_document.blocks:
        span = block.source_spans[0]
        assert source[span.char_start : span.char_end] == block.text
        assert span.line_start is not None and span.line_end is not None

    unknown_path = tmp_path / "notes.custom"
    unknown_path.write_text("Unknown suffix is plain text.", encoding="utf-8")
    unknown = parse_canonical_document(unknown_path)
    assert unknown.parser_source == "text"
    assert unknown.blocks[0].text == "Unknown suffix is plain text."

    empty_path = tmp_path / "empty.txt"
    empty_path.write_text("", encoding="utf-8")
    empty = parse_canonical_document(empty_path)
    assert empty.blocks == []
    assert empty.warnings


def test_text_source_ranges_preserve_crlf_offsets(tmp_path: Path) -> None:
    path = tmp_path / "windows-lines.txt"
    path.write_bytes(b"First paragraph.\r\n\r\nSecond paragraph.")
    with path.open("r", encoding="utf-8", newline="") as source_file:
        source = source_file.read()

    document = parse_canonical_document(path)
    assert len(document.blocks) == 2
    for block in document.blocks:
        span = block.source_spans[0]
        assert source[span.char_start : span.char_end] == block.text


def test_dispatcher_uses_pdf_adapter_without_recursing(monkeypatch, tmp_path: Path) -> None:
    pdf_path = tmp_path / "sample.pdf"
    pdf_path.write_bytes(b"placeholder")
    sentinel = object()
    monkeypatch.setattr(PDFCanonicalAdapter, "parse", lambda self, path: sentinel)
    assert parse_canonical_document(pdf_path) is sentinel


@pytest.mark.parametrize(
    ("suffix", "content", "parser_source"),
    [(".markdown", "# Alias", "markdown"), (".htm", "<p>Alias</p>", "html")],
)
def test_dispatcher_supports_format_aliases(
    tmp_path: Path, suffix: str, content: str, parser_source: str
) -> None:
    path = tmp_path / f"alias{suffix}"
    path.write_text(content, encoding="utf-8")
    assert parse_canonical_document(path).parser_source == parser_source


def test_missing_path_raises_clear_parse_error(tmp_path: Path) -> None:
    missing = tmp_path / "missing.md"
    with pytest.raises((DocumentParseError, ValueError), match="missing|exist|read"):
        parse_canonical_document(missing)


def test_read_failure_is_wrapped_as_document_parse_error(monkeypatch, tmp_path: Path) -> None:
    path = tmp_path / "unreadable.txt"
    path.write_text("content", encoding="utf-8")
    original = Path.open

    def fail_for_source(candidate: Path, mode="r", *args, **kwargs):
        if candidate == path.resolve() and "b" in mode:
            raise PermissionError("access denied")
        return original(candidate, mode, *args, **kwargs)

    monkeypatch.setattr(Path, "open", fail_for_source)
    with pytest.raises(DocumentParseError, match="Unable to read|access denied"):
        parse_canonical_document(path)


def test_compatibility_parse_document_preserves_long_canonical_block(tmp_path: Path) -> None:
    content = "x" * 5001
    path = tmp_path / "long.txt"
    path.write_text(content, encoding="utf-8")
    parsed = parse_document(path)
    assert len(parsed.chunks) == 1
    assert parsed.chunks[0].text == content
    assert parsed.text == content
    assert parsed.metadata["format"] == "text"
    assert parsed.metadata["canonical"]["parser_source"] == "text"
