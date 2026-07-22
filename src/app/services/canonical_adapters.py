from __future__ import annotations

import hashlib
import json
import mimetypes
import os
import re
import stat
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol
from urllib.parse import unquote, urlsplit
from uuid import uuid4

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from lxml import etree

from app.core.config import get_settings
from app.services.canonical_abstract import extract_explicit_abstract
from app.services.canonical_provenance import block_is_generated
from app.services.canonical_table_identity import (
    table_identity_fingerprint,
)
from app.services.canonical_models import (
    CanonicalAsset,
    CanonicalBlock,
    CanonicalCell,
    CanonicalDocument,
    CanonicalFigure,
    CanonicalFormula,
    CanonicalQualityIssue,
    CanonicalTable,
    SectionNode,
    SourceSpan,
)
from app.services.filesystem import display_title_from_path


MAX_TABLE_ROWS = 10_000
MAX_TABLE_COLUMNS = 1_000
MAX_TABLE_GRID_CELLS = 1_000_000
MAX_SINGLE_ASSET_BYTES = 64 * 1024 * 1024
MAX_DOCUMENT_ASSET_BYTES = 256 * 1024 * 1024
ASSET_COPY_CHUNK_BYTES = 1024 * 1024
ADAPTER_ASSET_CACHE_MAX_BYTES = 4 * 1024 * 1024 * 1024
ADAPTER_ASSET_CACHE_CLEANUP_POLICY = (
    "evict source-sha directories by oldest mtime, then lexical sha"
)
MAX_NESTING_DEPTH = 256


class CanonicalAdapter(Protocol):
    def parse(self, path: Path) -> CanonicalDocument: ...


@dataclass(frozen=True)
class _Line:
    number: int
    start: int
    content_end: int
    end: int
    text: str


@dataclass(frozen=True)
class _MarkdownImage:
    start: int
    end: int
    alt: str
    target: str
    title: str | None = None
    reference_label: str | None = None


def _parse_error(path: Path, message: str, exc: Exception | None = None) -> Exception:
    # Lazy import keeps parser.parse_document() free to import the dispatcher lazily.
    from app.services.parser import DocumentParseError

    error = DocumentParseError(path, message)
    if exc is not None:
        error.__cause__ = exc
    return error


def _validate_path(path: Path) -> Path:
    path = Path(path).expanduser()
    if not path.exists():
        raise _parse_error(path, "Source document does not exist.")
    if not path.is_file():
        raise _parse_error(path, "Source document is not a readable file.")
    if not os.access(path, os.R_OK):
        raise _parse_error(path, "Source document is not readable.")
    return path.resolve()


def _read_text(path: Path) -> str:
    try:
        with path.open("r", encoding="utf-8", newline="") as source_file:
            return source_file.read()
    except (OSError, UnicodeError) as exc:
        raise _parse_error(path, f"Unable to read text: {exc}", exc)


def _is_link_or_reparse_point(path: Path) -> bool:
    try:
        path_stat = path.lstat()
    except FileNotFoundError:
        return False
    reparse_attribute = getattr(stat, "FILE_ATTRIBUTE_REPARSE_POINT", 0x400)
    file_attributes = getattr(path_stat, "st_file_attributes", 0)
    return stat.S_ISLNK(path_stat.st_mode) or bool(file_attributes & reparse_attribute)


def _has_link_or_reparse_component(path: Path) -> bool:
    current = path
    while True:
        if _is_link_or_reparse_point(current):
            return True
        if current.parent == current:
            return False
        current = current.parent


def _safe_local_asset_name(path: Path, asset_sha: str) -> str:
    suffix = path.suffix.lower()
    if not re.fullmatch(r"\.[a-z0-9]{1,10}", suffix):
        suffix = ""
    stem = re.sub(r"[^A-Za-z0-9._-]+", "-", path.stem).strip(". -") or "asset"
    return f"{asset_sha[:24]}-{stem[:80]}{suffix}"


def _materialize_adapter_asset(
    document: CanonicalDocument,
    safe_name: str,
    blob: bytes | Path,
    expected_sha: str,
    *,
    source_kind: str,
    asset_cache_root: Path | None = None,
) -> Path:
    """Atomically persist an extracted asset outside parser temporary output."""

    source_path = Path(document.source_path or "document")
    try:
        if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9._-]*", safe_name):
            raise ValueError("asset cache filename is not portable")
        root = asset_cache_root
        if root is None:
            root = get_settings().cache_dir / "canonical_adapter_assets"
        root = Path(root).expanduser()
        if _is_link_or_reparse_point(root):
            raise ValueError("asset cache root is a symbolic link or reparse point")
        root.mkdir(parents=True, exist_ok=True)
        if _is_link_or_reparse_point(root):
            raise ValueError("asset cache root is a symbolic link or reparse point")
        resolved_root = root.resolve(strict=True)

        source_sha = str(document.source_metadata["sha256"])
        if not re.fullmatch(r"[0-9a-f]{64}", source_sha):
            raise ValueError("canonical source sha256 must be lowercase 64-hex")
        if not re.fullmatch(r"[0-9a-f]{64}", expected_sha):
            raise ValueError("canonical asset sha256 must be lowercase 64-hex")
        source_stat: os.stat_result | None = None
        if isinstance(blob, Path):
            if _has_link_or_reparse_component(blob):
                raise ValueError("extracted asset source contains a link or reparse point")
            source_stat = blob.stat()
            if source_stat.st_size > MAX_SINGLE_ASSET_BYTES:
                raise ValueError("extracted asset size limit exceeded")
            with blob.open("rb") as source:
                source_digest = hashlib.file_digest(source, "sha256").hexdigest()
        else:
            if len(blob) > MAX_SINGLE_ASSET_BYTES:
                raise ValueError("extracted asset size limit exceeded")
            source_digest = hashlib.sha256(blob).hexdigest()
        if source_digest != expected_sha:
            raise ValueError("extracted asset hash mismatch")

        directory = root / source_sha
        if _is_link_or_reparse_point(directory):
            raise ValueError(
                "asset cache source directory is a symbolic link or reparse point"
            )
        directory.mkdir(exist_ok=True)
        if _is_link_or_reparse_point(directory):
            raise ValueError(
                "asset cache source directory is a symbolic link or reparse point"
            )
        resolved_directory = directory.resolve(strict=True)
        try:
            resolved_directory.relative_to(resolved_root)
        except ValueError as exc:
            raise ValueError("asset cache source directory escapes cache root") from exc
        if resolved_directory.parent != resolved_root:
            raise ValueError("asset cache source directory is not directly under cache root")

        destination = directory / safe_name
        if _is_link_or_reparse_point(destination):
            raise ValueError(
                "asset cache destination is a symbolic link or reparse point"
            )
        if destination.parent.resolve(strict=True) != resolved_directory:
            raise ValueError("asset cache destination escapes source directory")
        if destination.is_file():
            with destination.open("rb") as persisted:
                persisted_sha = hashlib.file_digest(persisted, "sha256").hexdigest()
            if persisted_sha == expected_sha:
                return destination.resolve(strict=True)

        temporary = directory / f".{safe_name}.{uuid4().hex}.tmp"
        if _is_link_or_reparse_point(temporary):
            raise ValueError(
                "asset cache temporary path is a symbolic link or reparse point"
            )
        with temporary.open("xb") as output:
            copied_digest = hashlib.sha256()
            copied_size = 0
            if isinstance(blob, Path):
                with blob.open("rb") as source:
                    opened_stat = os.fstat(source.fileno())
                    if source_stat is None or (
                        opened_stat.st_size != source_stat.st_size
                        or opened_stat.st_mtime_ns != source_stat.st_mtime_ns
                    ):
                        raise ValueError("extracted asset changed before copy")
                    while chunk := source.read(ASSET_COPY_CHUNK_BYTES):
                        copied_size += len(chunk)
                        copied_digest.update(chunk)
                        output.write(chunk)
                    final_stat = os.fstat(source.fileno())
                if (
                    copied_size != source_stat.st_size
                    or final_stat.st_size != source_stat.st_size
                    or blob.stat().st_mtime_ns != source_stat.st_mtime_ns
                    or _has_link_or_reparse_component(blob)
                ):
                    raise ValueError("extracted asset changed during copy")
            else:
                copied_size = len(blob)
                copied_digest.update(blob)
                output.write(blob)
            output.flush()
            os.fsync(output.fileno())
        if copied_digest.hexdigest() != expected_sha:
            raise ValueError("materialized asset hash mismatch")
        if _is_link_or_reparse_point(temporary):
            raise ValueError(
                "asset cache temporary path became a symbolic link or reparse point"
            )
        with temporary.open("rb") as persisted:
            temporary_sha = hashlib.file_digest(persisted, "sha256").hexdigest()
        if temporary_sha != expected_sha:
            raise ValueError("materialized asset hash mismatch")
        if _is_link_or_reparse_point(destination):
            raise ValueError(
                "asset cache destination is a symbolic link or reparse point"
            )
        os.replace(temporary, destination)
        if _is_link_or_reparse_point(destination):
            raise ValueError(
                "asset cache destination became a symbolic link or reparse point"
            )
        with destination.open("rb") as persisted:
            destination_sha = hashlib.file_digest(persisted, "sha256").hexdigest()
        if destination_sha != expected_sha:
            raise ValueError("persisted asset hash mismatch")
        return destination.resolve(strict=True)
    except Exception as exc:
        raise _parse_error(
            source_path,
            f"Unable to materialize {source_kind} asset in cache: {exc}",
            exc,
        )
    finally:
        temporary_path = locals().get("temporary")
        if isinstance(temporary_path, Path):
            try:
                if temporary_path.exists() and not _is_link_or_reparse_point(
                    temporary_path
                ):
                    temporary_path.unlink()
            except OSError:
                pass


def _register_local_figure_asset(
    document: CanonicalDocument,
    raw_target: str | None,
    span: SourceSpan,
) -> str | None:
    target = (raw_target or "").strip()
    if not target:
        return None

    source_path = Path(document.source_path or "")
    try:
        split = urlsplit(target)
        scheme = split.scheme.lower()
        windows_absolute = bool(re.match(r"^[A-Za-z]:[\\/]", target))
        if split.netloc and scheme != "file":
            document.warnings.append(
                f"Figure target {target!r} is a non-local URI and was not materialized."
            )
            return None
        if scheme and not windows_absolute and scheme != "file":
            document.warnings.append(
                f"Figure target {target!r} is a non-local URI and was not materialized."
            )
            return None
        if scheme == "file":
            local_authority = not split.netloc
            if split.netloc:
                try:
                    hostname = split.hostname
                    port = split.port
                except ValueError:
                    hostname = None
                    port = None
                local_authority = bool(
                    hostname
                    and hostname.casefold() == "localhost"
                    and split.username is None
                    and split.password is None
                    and port is None
                )
            if not local_authority:
                document.warnings.append(
                    f"Figure target {target!r} is a non-local file URI and was not materialized."
                )
                return None
            decoded = unquote(split.path)
            if re.match(r"^/[A-Za-z]:/", decoded):
                decoded = decoded[1:]
            candidate = Path(decoded)
        elif windows_absolute:
            candidate = Path(unquote(target.split("?", 1)[0].split("#", 1)[0]))
        else:
            candidate = Path(unquote(split.path))
            if not candidate.is_absolute():
                candidate = source_path.parent / candidate

        if (
            not candidate.exists()
            or not candidate.is_file()
            or _has_link_or_reparse_component(candidate)
        ):
            document.warnings.append(
                f"Figure target {target!r} is missing, not a regular file, or is a link; "
                "it was not materialized."
            )
            return None
        resolved = candidate.resolve(strict=True)
        size_bytes = resolved.stat().st_size
        existing_total = sum(
            int(asset.metadata.get("size_bytes") or 0) for asset in document.assets
        )
        if (
            size_bytes > MAX_SINGLE_ASSET_BYTES
            or existing_total + size_bytes > MAX_DOCUMENT_ASSET_BYTES
        ):
            raise ValueError("canonical asset size limit exceeded")
        with resolved.open("rb") as source:
            asset_sha = hashlib.file_digest(source, "sha256").hexdigest()
    except (OSError, ValueError) as exc:
        document.warnings.append(
            f"Figure target {target!r} could not be read and was not materialized: {exc}."
        )
        return None

    existing = next((asset for asset in document.assets if asset.sha256 == asset_sha), None)
    if existing is not None:
        if span not in existing.source_spans:
            existing.source_spans.append(span)
        return existing.path

    safe_name = _safe_local_asset_name(resolved, asset_sha)
    asset_path = f"assets/{safe_name}"
    media_type = mimetypes.guess_type(resolved.name)[0] or "application/octet-stream"
    document.assets.append(
        CanonicalAsset(
            asset_id=_stable_id("asset", document.document_id, asset_sha),
            path=asset_path,
            media_type=media_type,
            sha256=asset_sha,
            source_path=str(resolved),
            source_spans=[span],
            metadata={"source_target": target, "size_bytes": size_bytes},
        )
    )
    return asset_path


def _raise_table_limit(document: CanonicalDocument, source_format: str, detail: str) -> None:
    raise _parse_error(
        Path(document.source_path or f"document.{source_format.lower()}"),
        f"{source_format} table exceeds canonical {detail} limit.",
    )


def _validate_html_nesting_depth(root: Tag, path: Path) -> None:
    stack: list[tuple[Tag, int]] = [(root, 0)]
    while stack:
        node, depth = stack.pop()
        if depth > MAX_NESTING_DEPTH:
            raise _parse_error(
                path,
                f"HTML nesting depth exceeds the {MAX_NESTING_DEPTH} element limit.",
            )
        children = [child for child in node.children if isinstance(child, Tag)]
        stack.extend((child, depth + 1) for child in reversed(children))


def _stable_id(prefix: str, *parts: object) -> str:
    payload = "\x1f".join(str(part) for part in parts)
    return f"{prefix}-{hashlib.sha256(payload.encode('utf-8')).hexdigest()[:20]}"


def _new_document(path: Path, parser_source: str, media_type: str) -> CanonicalDocument:
    try:
        stat = path.stat()
        with path.open("rb") as source_file:
            source_hash = hashlib.file_digest(source_file, "sha256").hexdigest()
    except OSError as exc:
        raise _parse_error(path, f"Unable to read source document: {exc}", exc)
    return CanonicalDocument(
        document_id=_stable_id("doc", parser_source, source_hash),
        source_path=str(path),
        source_media_type=media_type,
        parser_source=parser_source,
        parse_version="canonical-v1",
        title=display_title_from_path(path),
        source_metadata={
            "filename": path.name,
            "size_bytes": stat.st_size,
            "sha256": source_hash,
        },
        parser_metadata={
            "adapter": parser_source,
            "format": parser_source,
            "canonical_schema": "v1",
        },
        metadata={"source_suffix": path.suffix.lower()},
    )


class _DocumentBuilder:
    def __init__(self, document: CanonicalDocument) -> None:
        self.document = document
        self._headings: list[tuple[int, str]] = []
        self._outline_stack: list[tuple[int, SectionNode]] = []
        self._reference_level: int | None = None
        self._appendix_level: int | None = None

    @property
    def heading_path(self) -> list[str]:
        return [title for _, title in self._headings]

    def _span(self, span: SourceSpan) -> SourceSpan:
        return span.model_copy(update={"heading_path": self.heading_path})

    def _block_id(self, block_type: str, span: SourceSpan, text: str) -> str:
        locator = json.dumps(span.model_dump(mode="json"), sort_keys=True, ensure_ascii=True)
        return _stable_id(
            "block",
            self.document.document_id,
            len(self.document.blocks),
            block_type,
            locator,
            text,
        )

    def add_heading(self, title: str, level: int, span: SourceSpan, **metadata: object) -> CanonicalBlock:
        if self._reference_level is not None and level <= self._reference_level:
            self._reference_level = None
        if self._appendix_level is not None and level <= self._appendix_level:
            self._appendix_level = None

        while self._headings and self._headings[-1][0] >= level:
            self._headings.pop()
        self._headings.append((level, title))

        normalized = title.strip().lower().rstrip(":")
        if (
            normalized in {"references", "reference", "bibliography"}
            and self._reference_level is None
        ):
            self._reference_level = level
            self._appendix_level = None
        elif normalized.startswith("appendix") and self._reference_level is None:
            self._appendix_level = level

        updated_span = self._span(span)
        block = self._append_block(
            block_type="heading",
            text=title,
            span=updated_span,
            metadata={"heading_level": level, **metadata},
        )

        node = SectionNode(title=title, level=level, block_id=block.block_id)
        while self._outline_stack and self._outline_stack[-1][0] >= level:
            self._outline_stack.pop()
        if self._outline_stack:
            self._outline_stack[-1][1].children.append(node)
        else:
            self.document.outline.append(node)
        self._outline_stack.append((level, node))
        return block

    def add_content(
        self,
        text: str,
        span: SourceSpan,
        *,
        block_type: str = "narrative",
        metadata: dict[str, object] | None = None,
        table_id: str | None = None,
        figure_id: str | None = None,
        formula_id: str | None = None,
        retrievable: bool | None = None,
    ) -> CanonicalBlock:
        if block_type == "narrative":
            if self._reference_level is not None:
                block_type = "reference"
            elif self._appendix_level is not None:
                block_type = "appendix"
        return self._append_block(
            block_type=block_type,
            text=text,
            span=self._span(span),
            metadata=metadata or {},
            table_id=table_id,
            figure_id=figure_id,
            formula_id=formula_id,
            retrievable=retrievable,
        )

    def _append_block(
        self,
        *,
        block_type: str,
        text: str,
        span: SourceSpan,
        metadata: dict[str, object],
        table_id: str | None = None,
        figure_id: str | None = None,
        formula_id: str | None = None,
        retrievable: bool | None = None,
    ) -> CanonicalBlock:
        default_retrievable = block_type != "heading" and self._reference_level is None
        block = CanonicalBlock(
            block_id=self._block_id(block_type, span, text),
            block_type=block_type,
            text=text,
            section_path=self.heading_path,
            reading_order=len(self.document.blocks),
            source_spans=[span],
            parser_source=self.document.parser_source,
            table_id=table_id,
            figure_id=figure_id,
            formula_id=formula_id,
            metadata=metadata,
            retrievable=default_retrievable and retrievable is not False,
        )
        self.document.blocks.append(block)
        return block


def _lines(source: str) -> list[_Line]:
    records: list[_Line] = []
    offset = 0
    for number, raw in enumerate(source.splitlines(keepends=True)):
        raw_content = raw.rstrip("\r\n")
        has_bom = number == 0 and raw_content.startswith("\ufeff")
        content = raw_content[1:] if has_bom else raw_content
        records.append(
            _Line(
                number=number,
                start=offset + (1 if has_bom else 0),
                content_end=offset + len(raw_content),
                end=offset + len(raw),
                text=content,
            )
        )
        offset += len(raw)
    if source and not records:
        records.append(_Line(0, 0, len(source), len(source), source))
    return records


def _text_span(lines: list[_Line], start: int, end: int) -> SourceSpan:
    return SourceSpan(
        line_start=lines[start].number,
        line_end=lines[end].number,
        char_start=lines[start].start,
        char_end=lines[end].content_end,
    )


def _span_for_chars(lines: list[_Line], char_start: int, char_end: int) -> SourceSpan:
    start_line = next(
        line for line in lines if line.start <= char_start < max(line.end, line.start + 1)
    )
    end_position = max(char_start, char_end - 1)
    end_line = next(
        line for line in lines if line.start <= end_position < max(line.end, line.start + 1)
    )
    return SourceSpan(
        line_start=start_line.number,
        line_end=end_line.number,
        char_start=char_start,
        char_end=char_end,
    )


def _table_markdown(headers: list[str], rows: list[list[str]]) -> str:
    width = max([len(headers), *(len(row) for row in rows)], default=0)
    if width == 0:
        return ""
    padded_headers = (headers + [""] * width)[:width]
    lines = [
        "| " + " | ".join(_escape_markdown_cell(cell) for cell in padded_headers) + " |",
        "| " + " | ".join("---" for _ in range(width)) + " |",
    ]
    for row in rows:
        padded = (row + [""] * width)[:width]
        lines.append("| " + " | ".join(_escape_markdown_cell(cell) for cell in padded) + " |")
    return "\n".join(lines)


def _escape_markdown_cell(value: str) -> str:
    return value.replace("\\", "\\\\").replace("|", "\\|").replace("\n", "<br>")


def _split_pipe_row(value: str) -> list[str]:
    value = value.strip()
    if value.startswith("|"):
        value = value[1:]
    if value.endswith("|") and not value.endswith("\\|"):
        value = value[:-1]
    parts = re.split(r"(?<!\\)\|", value)
    return [part.strip().replace("\\|", "|") for part in parts]


def _is_table_separator(value: str) -> bool:
    cells = _split_pipe_row(value)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


class TextCanonicalAdapter:
    parser_source = "text"

    def parse(self, path: Path) -> CanonicalDocument:
        path = _validate_path(path)
        source = _read_text(path)
        media_type = mimetypes.guess_type(path.name)[0] or "text/plain"
        document = _new_document(path, self.parser_source, media_type)
        builder = _DocumentBuilder(document)
        lines = _lines(source)

        index = 0
        while index < len(lines):
            if not lines[index].text.strip():
                index += 1
                continue
            start = index
            while index + 1 < len(lines) and lines[index + 1].text.strip():
                index += 1
            end = index
            span = _text_span(lines, start, end)
            builder.add_content(source[span.char_start : span.char_end], span)
            index += 1

        if not document.blocks:
            document.warnings.append("Source text is empty; canonical document contains no blocks.")
        return document


class MarkdownCanonicalAdapter:
    parser_source = "markdown"
    _heading_re = re.compile(
        r"^[ ]{0,3}(#{1,6})(?P<content>(?:[ \t]+.*)?)$"
    )
    _fence_re = re.compile(r"^[ \t]*(`{3,}|~{3,})(.*)$")
    _setext_re = re.compile(r"^[ \t]*(?P<underline>=+|-+)[ \t]*$")
    _definition_re = re.compile(
        r"^[ ]{0,3}\[(?P<label>[^]]+)\]:[ \t]*"
        r"(?:<(?P<angle>[^>]*)>|(?P<bare>\S+))"
        r"(?:[ \t]+(?:\"(?P<double_title>[^\"]*)\"|'(?P<single_title>[^']*)'|"
        r"\((?P<paren_title>[^)]*)\)))?[ \t]*$"
    )

    def parse(self, path: Path) -> CanonicalDocument:
        path = _validate_path(path)
        source = _read_text(path)
        document = _new_document(path, self.parser_source, "text/markdown")
        builder = _DocumentBuilder(document)
        lines = _lines(source)
        definitions, definition_lines = self._definitions(lines)
        index = 0

        while index < len(lines):
            if index in definition_lines:
                index += 1
                continue
            if not lines[index].text.strip():
                index += 1
                continue

            heading_match = self._heading_re.match(lines[index].text)
            if heading_match:
                title = self._atx_heading_title(heading_match.group("content"))
                builder.add_heading(title, len(heading_match.group(1)), _text_span(lines, index, index))
                if not document.title or document.title == display_title_from_path(path):
                    if title and len(heading_match.group(1)) == 1:
                        document.title = title
                index += 1
                continue

            setext_level = self._setext_level(lines, index)
            if setext_level is not None:
                title = lines[index].text.strip()
                builder.add_heading(title, setext_level, _text_span(lines, index, index + 1))
                if setext_level == 1 and document.title == display_title_from_path(path):
                    document.title = title
                index += 2
                continue

            fence_match = self._fence_re.match(lines[index].text)
            if fence_match:
                fence = fence_match.group(1)
                end = index + 1
                close_re = re.compile(rf"^[ \t]*{re.escape(fence[0])}{{{len(fence)},}}[ \t]*$")
                while end < len(lines) and not close_re.match(lines[end].text):
                    end += 1
                closed = end < len(lines)
                last = end if closed else len(lines) - 1
                body_start = lines[index].end
                body_end = lines[end].start if closed else len(source)
                body = source[body_start:body_end]
                code_span = _text_span(lines, index, last)
                if not closed:
                    code_span = code_span.model_copy(update={"char_end": len(source)})
                info = fence_match.group(2).strip()
                builder.add_content(
                    body,
                    code_span,
                    metadata={
                        "kind": "code",
                        "language": info.split(maxsplit=1)[0] if info else "",
                        "info": info,
                        "fence": fence,
                        "closed": closed,
                    },
                )
                if not closed:
                    document.warnings.append(f"Unclosed Markdown code fence at line {index + 1}.")
                index = last + 1
                continue

            formula = self._formula_at(source, lines, index)
            if formula is not None:
                latex, end, delimiter = formula
                span = _text_span(lines, index, end)
                formula_id = _stable_id("formula", document.document_id, span.char_start, latex)
                canonical_formula = CanonicalFormula(
                    formula_id=formula_id,
                    latex=latex,
                    source_spans=[span.model_copy(update={"heading_path": builder.heading_path})],
                    metadata={"delimiter": delimiter, "source_format": "latex"},
                )
                document.formulas.append(canonical_formula)
                builder.add_content(
                    latex,
                    span,
                    block_type="formula",
                    formula_id=formula_id,
                    metadata={"delimiter": delimiter},
                )
                index = end + 1
                continue

            if index + 1 < len(lines) and "|" in lines[index].text and _is_table_separator(lines[index + 1].text):
                end = index + 2
                while end < len(lines) and lines[end].text.strip() and "|" in lines[end].text:
                    end += 1
                last = end - 1
                headers = _split_pipe_row(lines[index].text)
                rows = [_split_pipe_row(lines[row].text) for row in range(index + 2, end)]
                span = _text_span(lines, index, last)
                source_markdown = source[span.char_start : span.char_end]
                table_id = _stable_id(
                    "table", document.document_id, span.char_start, source_markdown
                )
                cells = [
                    CanonicalCell(text=value, row_index=0, column_index=column, is_header=True)
                    for column, value in enumerate(headers)
                ]
                cells.extend(
                    CanonicalCell(text=value, row_index=row_index, column_index=column)
                    for row_index, row in enumerate(rows, start=1)
                    for column, value in enumerate(row)
                )
                table = CanonicalTable(
                    table_id=table_id,
                    headers=headers,
                    rows=rows,
                    cells=cells,
                    source_markdown=source_markdown,
                    normalized_markdown=_table_markdown(headers, rows),
                    source_spans=[span.model_copy(update={"heading_path": builder.heading_path})],
                )
                document.tables.append(table)
                builder.add_content(
                    table.normalized_markdown or source_markdown,
                    span,
                    block_type="table",
                    table_id=table_id,
                )
                index = end
                continue

            start = index
            index += 1
            while (
                index < len(lines)
                and index not in definition_lines
                and lines[index].text.strip()
                and not self._is_special(lines, index)
            ):
                index += 1
            end = index - 1
            span = _text_span(lines, start, end)
            text = source[span.char_start : span.char_end]
            blocks = self._add_paragraph(
                document, builder, lines, text, span, definitions
            )
            if builder.heading_path and builder.heading_path[-1].strip().lower().rstrip(":") == "abstract":
                abstract_text = "".join(
                    block.text for block in blocks if block.block_type == "narrative"
                )
                document.abstract = document.abstract or abstract_text or None

        if not document.blocks:
            document.warnings.append("Source Markdown is empty; canonical document contains no blocks.")
        return document

    def _is_special(self, lines: list[_Line], index: int) -> bool:
        value = lines[index].text
        return bool(
            self._heading_re.match(value)
            or self._fence_re.match(value)
            or self._setext_level(lines, index) is not None
            or value.strip().startswith(("$$", "\\["))
            or (
                index + 1 < len(lines)
                and "|" in value
                and _is_table_separator(lines[index + 1].text)
            )
        )

    def _setext_level(self, lines: list[_Line], index: int) -> int | None:
        if index + 1 >= len(lines) or not lines[index].text.strip():
            return None
        match = self._setext_re.fullmatch(lines[index + 1].text)
        if match is None:
            return None
        return 1 if match.group("underline").startswith("=") else 2

    @staticmethod
    def _atx_heading_title(raw_content: str) -> str:
        closing = re.fullmatch(
            r"(?P<title>.*?)[ \t]+#+[ \t]*",
            raw_content,
        )
        if closing is not None:
            raw_content = closing.group("title")
        return raw_content.strip(" \t")

    def _add_paragraph(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        lines: list[_Line],
        text: str,
        span: SourceSpan,
        definitions: dict[str, tuple[str, str | None]],
    ) -> list[CanonicalBlock]:
        blocks: list[CanonicalBlock] = []
        absolute_start = span.char_start or 0
        cursor = 0
        for image in self._scan_images(text, definitions):
            if image.start > cursor:
                before = text[cursor : image.start]
                if before.strip():
                    before_start = absolute_start + cursor
                    blocks.append(
                        builder.add_content(
                            before,
                            _span_for_chars(
                                lines, before_start, absolute_start + image.start
                            ),
                        )
                    )
            raw = text[image.start : image.end]
            image_start = absolute_start + image.start
            blocks.append(
                self._add_image(
                    document,
                    builder,
                    _span_for_chars(lines, image_start, image_start + len(raw)),
                    raw,
                    image,
                )
            )
            cursor = image.end
        if cursor < len(text):
            after = text[cursor:]
            if after.strip():
                after_start = absolute_start + cursor
                blocks.append(
                    builder.add_content(
                        after,
                        _span_for_chars(lines, after_start, absolute_start + len(text)),
                    )
                )
        if not blocks and text.strip():
            blocks.append(builder.add_content(text, span))
        return blocks

    @staticmethod
    def _add_image(
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        span: SourceSpan,
        raw: str,
        image: _MarkdownImage,
    ) -> CanonicalBlock:
        target = image.target
        title = image.title
        caption = image.alt or title or None
        figure_span = span.model_copy(update={"heading_path": builder.heading_path})
        asset_path = _register_local_figure_asset(document, target, figure_span)
        figure_id = _stable_id(
            "figure", document.document_id, span.char_start, target, caption
        )
        document.figures.append(
            CanonicalFigure(
                figure_id=figure_id,
                caption=caption,
                asset_path=asset_path,
                source_spans=[figure_span],
                metadata={
                    "target": target,
                    "alt": image.alt,
                    "title": title,
                    "reference_label": image.reference_label,
                    "source_markdown": raw,
                },
            )
        )
        return builder.add_content(
            caption or target,
            span,
            block_type="figure",
            figure_id=figure_id,
            metadata={"source_markdown": raw, "target": target, "title": title},
        )

    @classmethod
    def _definitions(
        cls, lines: list[_Line]
    ) -> tuple[dict[str, tuple[str, str | None]], set[int]]:
        definitions: dict[str, tuple[str, str | None]] = {}
        definition_lines: set[int] = set()
        active_fence: str | None = None
        for index, line in enumerate(lines):
            if active_fence is not None:
                close_re = re.compile(
                    rf"^[ \t]*{re.escape(active_fence[0])}{{{len(active_fence)},}}[ \t]*$"
                )
                if close_re.match(line.text):
                    active_fence = None
                continue
            fence_match = cls._fence_re.match(line.text)
            if fence_match is not None:
                active_fence = fence_match.group(1)
                continue
            match = cls._definition_re.fullmatch(line.text)
            if match is None:
                continue
            label = cls._normalize_reference_label(match.group("label"))
            angle_target = match.group("angle")
            target = angle_target if angle_target is not None else match.group("bare")
            title = (
                match.group("double_title")
                or match.group("single_title")
                or match.group("paren_title")
            )
            definitions.setdefault(
                label, (cls._unescape_markdown(target), title)
            )
            definition_lines.add(index)
        return definitions, definition_lines

    @classmethod
    def _scan_images(
        cls,
        text: str,
        definitions: dict[str, tuple[str, str | None]],
    ) -> list[_MarkdownImage]:
        images: list[_MarkdownImage] = []
        index = 0
        while index < len(text):
            if text[index] == "`":
                run_length = len(text[index:]) - len(text[index:].lstrip("`"))
                closing = cls._find_backtick_closing(
                    text, index + run_length, run_length
                )
                index = closing + run_length if closing >= 0 else index + run_length
                continue
            if (
                text[index] != "!"
                or cls._is_escaped(text, index)
                or index + 1 >= len(text)
                or text[index + 1] != "["
            ):
                index += 1
                continue
            alt_end = cls._find_closing_bracket(text, index + 2)
            if alt_end < 0:
                index += 1
                continue
            alt = cls._unescape_markdown(text[index + 2 : alt_end])
            cursor = alt_end + 1
            parsed: tuple[int, str, str | None, str | None] | None = None
            if cursor < len(text) and text[cursor] == "(":
                direct = cls._parse_direct_image(text, cursor)
                if direct is not None:
                    end, target, title = direct
                    parsed = (end, target, title, None)
            elif cursor < len(text) and text[cursor] == "[":
                label_end = cls._find_closing_bracket(text, cursor + 1)
                if label_end >= 0:
                    raw_label = text[cursor + 1 : label_end] or alt
                    label = cls._normalize_reference_label(raw_label)
                    if label in definitions:
                        target, title = definitions[label]
                        parsed = (label_end + 1, target, title, label)
            else:
                label = cls._normalize_reference_label(alt)
                if label in definitions:
                    target, title = definitions[label]
                    parsed = (cursor, target, title, label)
            if parsed is None:
                index += 1
                continue
            end, target, title, reference_label = parsed
            images.append(
                _MarkdownImage(
                    start=index,
                    end=end,
                    alt=alt,
                    target=target,
                    title=title,
                    reference_label=reference_label,
                )
            )
            index = end
        return images

    @classmethod
    def _parse_direct_image(
        cls, text: str, opening: int
    ) -> tuple[int, str, str | None] | None:
        cursor = opening + 1
        while cursor < len(text) and text[cursor] in " \t\r\n":
            cursor += 1
        if cursor >= len(text):
            return None
        if text[cursor] == ")":
            return cursor + 1, "", None
        if text[cursor] == "<":
            target_end = cls._find_unescaped(text, ">", cursor + 1)
            if target_end < 0:
                return None
            target = text[cursor + 1 : target_end]
            cursor = target_end + 1
        else:
            target_start = cursor
            depth = 0
            while cursor < len(text):
                character = text[cursor]
                if character == "\\" and cursor + 1 < len(text):
                    cursor += 2
                    continue
                if character == "(":
                    depth += 1
                elif character == ")":
                    if depth == 0:
                        break
                    depth -= 1
                elif character in " \t\r\n" and depth == 0:
                    break
                cursor += 1
            target = text[target_start:cursor]
            if not target:
                return None
        target = cls._unescape_markdown(target)
        while cursor < len(text) and text[cursor] in " \t\r\n":
            cursor += 1
        title: str | None = None
        if cursor < len(text) and text[cursor] in {'"', "'", "("}:
            opener = text[cursor]
            closer = ")" if opener == "(" else opener
            title_end = cls._find_unescaped(text, closer, cursor + 1)
            if title_end < 0:
                return None
            title = cls._unescape_markdown(text[cursor + 1 : title_end])
            cursor = title_end + 1
            while cursor < len(text) and text[cursor] in " \t\r\n":
                cursor += 1
        if cursor >= len(text) or text[cursor] != ")":
            return None
        return cursor + 1, target, title

    @staticmethod
    def _find_unescaped(text: str, character: str, start: int) -> int:
        index = start
        while index < len(text):
            if text[index] == "\\":
                index += 2
                continue
            if text[index] == character:
                return index
            index += 1
        return -1

    @staticmethod
    def _find_closing_bracket(text: str, start: int) -> int:
        depth = 0
        index = start
        while index < len(text):
            if text[index] == "\\":
                index += 2
                continue
            if text[index] == "[":
                depth += 1
            elif text[index] == "]":
                if depth == 0:
                    return index
                depth -= 1
            index += 1
        return -1

    @staticmethod
    def _find_backtick_closing(text: str, start: int, run_length: int) -> int:
        index = start
        while index < len(text):
            if text[index] != "`":
                index += 1
                continue
            candidate_length = len(text[index:]) - len(text[index:].lstrip("`"))
            if candidate_length == run_length:
                return index
            index += candidate_length
        return -1

    @staticmethod
    def _is_escaped(text: str, index: int) -> bool:
        slash_count = 0
        cursor = index - 1
        while cursor >= 0 and text[cursor] == "\\":
            slash_count += 1
            cursor -= 1
        return slash_count % 2 == 1

    @staticmethod
    def _normalize_reference_label(label: str) -> str:
        return re.sub(r"\s+", " ", label.strip()).casefold()

    @staticmethod
    def _unescape_markdown(value: str) -> str:
        return re.sub(r"\\([!\"#$%&'()*+,\-./:;<=>?@\[\\\]^_`{|}~])", r"\1", value)

    @staticmethod
    def _formula_at(source: str, lines: list[_Line], index: int) -> tuple[str, int, str] | None:
        stripped = lines[index].text.strip()
        if stripped.startswith("$$"):
            delimiter, closing = "$$", "$$"
        elif stripped.startswith("\\["):
            delimiter, closing = "\\[", "\\]"
        else:
            return None

        first = stripped[len(delimiter) :]
        if first.endswith(closing) and len(first) >= len(closing):
            return first[: -len(closing)].strip(), index, delimiter

        end = index + 1
        while end < len(lines) and not lines[end].text.strip().endswith(closing):
            end += 1
        if end >= len(lines):
            end = len(lines) - 1
            latex_end = lines[end].content_end
        else:
            latex_end = lines[end].content_end - len(closing)
        latex_start = lines[index].start + lines[index].text.find(delimiter) + len(delimiter)
        return source[latex_start:latex_end].strip(), end, delimiter


def _html_selector(element: Tag) -> str:
    if element.get("id"):
        return f"#{element['id']}"
    parts: list[str] = []
    current: Tag | None = element
    while current is not None and current.name != "[document]":
        siblings = (
            current.parent.find_all(current.name, recursive=False) if current.parent else []
        )
        position = next(
            (index for index, sibling in enumerate(siblings, start=1) if sibling is current),
            1,
        )
        parts.append(f"{current.name}:nth-of-type({position})")
        current = current.parent if isinstance(current.parent, Tag) else None
    return " > ".join(reversed(parts))


def _html_xpath(element: Tag) -> str:
    parts: list[str] = []
    current: Tag | None = element
    while current is not None and current.name != "[document]":
        siblings = (
            current.parent.find_all(current.name, recursive=False) if current.parent else []
        )
        position = next(
            (index for index, sibling in enumerate(siblings, start=1) if sibling is current),
            1,
        )
        parts.append(f"{current.name}[{position}]")
        current = current.parent if isinstance(current.parent, Tag) else None
    return "/" + "/".join(reversed(parts))


def _html_span(element: Tag) -> SourceSpan:
    return SourceSpan(
        xpath=_html_xpath(element),
        css_selector=_html_selector(element),
        element_id=element.get("id"),
    )


def _is_math_element(element: Tag) -> bool:
    classes = " ".join(element.get("class", [])).lower()
    return element.name == "math" or any(token in classes for token in ("math", "latex", "equation"))


class HtmlCanonicalAdapter:
    parser_source = "html"

    def parse(self, path: Path) -> CanonicalDocument:
        path = _validate_path(path)
        source = _read_text(path)
        soup = BeautifulSoup(source, "html.parser")
        _validate_html_nesting_depth(soup, path)
        html_title = soup.title.get_text(" ", strip=True) if soup.title else ""
        for unwanted in soup.find_all(["script", "style", "nav", "noscript", "template"]):
            unwanted.decompose()

        document = _new_document(path, self.parser_source, "text/html")
        builder = _DocumentBuilder(document)
        first_h1 = soup.find("h1")
        document.title = html_title or (
            first_h1.get_text(" ", strip=True) if first_h1 else display_title_from_path(path)
        )

        root = soup.body or soup
        for child in list(root.children):
            self._walk_node(document, builder, child)

        if not document.blocks:
            document.warnings.append("HTML contains no supported content blocks.")
        return document

    def _walk_node(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        node: Tag | NavigableString,
    ) -> None:
        if isinstance(node, NavigableString):
            if str(node).strip() and isinstance(node.parent, Tag):
                builder.add_content(str(node), _html_span(node.parent))
            return
        if not isinstance(node, Tag):
            return
        if re.fullmatch(r"h[1-6]", node.name):
            builder.add_heading(node.get_text(" ", strip=True), int(node.name[1]), _html_span(node))
            return
        if node.name == "table":
            self._walk_table_tree(document, builder, node)
            return
        if node.name == "figure":
            self._add_figure(document, builder, node)
            return
        if node.name == "p":
            self._walk_inline_container(document, builder, node)
            return
        if node.name == "pre":
            builder.add_content(
                node.get_text("", strip=False).strip("\r\n"),
                _html_span(node),
                metadata={"kind": "code", "language": self._code_language(node.find("code"))},
            )
            return
        if node.name == "code":
            builder.add_content(
                node.get_text("", strip=False),
                _html_span(node),
                metadata={"kind": "code", "language": self._code_language(node)},
            )
            return
        if _is_math_element(node):
            self._add_formula(document, builder, node)
            return
        if node.name == "img":
            self._add_image(document, builder, node)
            return
        for child in list(node.children):
            self._walk_node(document, builder, child)

    def _walk_table_tree(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        table: Tag,
        *,
        parent_table_id: str | None = None,
        parent_row_index: int | None = None,
        parent_column_index: int | None = None,
        parent_figure_id: str | None = None,
    ) -> None:
        table_id, cell_positions = self._add_table(
            document,
            builder,
            table,
            parent_table_id=parent_table_id,
            parent_row_index=parent_row_index,
            parent_column_index=parent_column_index,
            parent_figure_id=parent_figure_id,
        )
        for cell, row_index, column_index in cell_positions:
            self._walk_structured_children(
                document,
                builder,
                cell,
                parent_table_id=table_id,
                parent_row_index=row_index,
                parent_column_index=column_index,
                parent_figure_id=parent_figure_id,
            )

    def _walk_structured_children(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        container: Tag,
        *,
        parent_table_id: str | None = None,
        parent_row_index: int | None = None,
        parent_column_index: int | None = None,
        parent_figure_id: str | None = None,
        skip_image: Tag | None = None,
    ) -> None:
        context = self._context_metadata(
            parent_table_id,
            parent_row_index,
            parent_column_index,
            parent_figure_id,
        )
        covered_context = self._with_table_retrieval_coverage(context)

        def visit(node: Tag | NavigableString) -> None:
            if not isinstance(node, Tag):
                return
            span = self._context_span(node, context)
            if node.name == "pre":
                builder.add_content(
                    node.get_text("", strip=False).strip("\r\n"),
                    span,
                    metadata={
                        "kind": "code",
                        "language": self._code_language(node.find("code")),
                        **covered_context,
                    },
                    retrievable=(
                        False
                        if covered_context.get("retrieval_covered_by_table_id")
                        else None
                    ),
                )
                return
            if node.name == "code":
                builder.add_content(
                    node.get_text("", strip=False),
                    span,
                    metadata={
                        "kind": "code",
                        "language": self._code_language(node),
                        **covered_context,
                    },
                    retrievable=(
                        False
                        if covered_context.get("retrieval_covered_by_table_id")
                        else None
                    ),
                )
                return
            if _is_math_element(node):
                self._add_formula(document, builder, node, context=covered_context)
                return
            if node.name == "img":
                if node is not skip_image:
                    self._add_image(document, builder, node, context=covered_context)
                return
            if node.name == "table":
                self._walk_table_tree(
                    document,
                    builder,
                    node,
                    parent_table_id=parent_table_id,
                    parent_row_index=parent_row_index,
                    parent_column_index=parent_column_index,
                    parent_figure_id=parent_figure_id,
                )
                return
            if node.name == "figure":
                self._add_figure(document, builder, node, context=covered_context)
                return
            if node.name == "figcaption":
                return
            for child in list(node.children):
                visit(child)

        for child in list(container.children):
            visit(child)

    @staticmethod
    def _context_metadata(
        parent_table_id: str | None,
        parent_row_index: int | None,
        parent_column_index: int | None,
        parent_figure_id: str | None,
    ) -> dict[str, object]:
        values = {
            "parent_table_id": parent_table_id,
            "parent_row_index": parent_row_index,
            "parent_column_index": parent_column_index,
            "parent_figure_id": parent_figure_id,
        }
        return {key: value for key, value in values.items() if value is not None}

    @staticmethod
    def _with_table_retrieval_coverage(
        context: dict[str, object],
    ) -> dict[str, object]:
        table_id = context.get("parent_table_id")
        if table_id is None:
            return context
        return {**context, "retrieval_covered_by_table_id": table_id}

    @staticmethod
    def _context_span(element: Tag, context: dict[str, object]) -> SourceSpan:
        span = _html_span(element)
        return span.model_copy(
            update={
                "table_id": context.get("parent_table_id"),
                "row_index": context.get("parent_row_index"),
                "column_index": context.get("parent_column_index"),
                "metadata": dict(context),
            }
        )

    def _walk_inline_container(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        container: Tag,
        *,
        context: dict[str, object] | None = None,
    ) -> None:
        context = dict(context or {})
        buffer: list[str] = []
        text_index = 0

        def flush() -> None:
            nonlocal text_index
            text = "".join(buffer)
            buffer.clear()
            if not text.strip():
                return
            parent_span = self._context_span(container, context)
            builder.add_content(
                text,
                parent_span.model_copy(
                    update={
                        "metadata": {**context, "segment_index": text_index},
                    }
                ),
                metadata=dict(context),
                retrievable=(
                    False if context.get("retrieval_covered_by_table_id") else None
                ),
            )
            text_index += 1

        def visit(node: Tag | NavigableString) -> None:
            if isinstance(node, NavigableString):
                buffer.append(str(node))
                return
            if not isinstance(node, Tag):
                return
            if node.name == "br":
                buffer.append("\n")
                return
            if node.name == "code":
                flush()
                structure_context = self._with_table_retrieval_coverage(context)
                builder.add_content(
                    node.get_text("", strip=False),
                    self._context_span(node, structure_context),
                    metadata={
                        "kind": "code",
                        "language": self._code_language(node),
                        **structure_context,
                    },
                    retrievable=(
                        False
                        if structure_context.get("retrieval_covered_by_table_id")
                        else None
                    ),
                )
                return
            if _is_math_element(node):
                flush()
                self._add_formula(
                    document,
                    builder,
                    node,
                    context=self._with_table_retrieval_coverage(context),
                )
                return
            if node.name == "img":
                flush()
                self._add_image(
                    document,
                    builder,
                    node,
                    context=self._with_table_retrieval_coverage(context),
                )
                return
            for child in list(node.children):
                visit(child)

        for child in list(container.children):
            visit(child)
        flush()

    @staticmethod
    def _code_language(element: Tag | None) -> str:
        if element is None:
            return ""
        for class_name in element.get("class", []):
            if class_name.startswith("language-"):
                return class_name.removeprefix("language-")
        return ""

    @classmethod
    def _add_table(
        cls,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        element: Tag,
        *,
        parent_table_id: str | None = None,
        parent_row_index: int | None = None,
        parent_column_index: int | None = None,
        parent_figure_id: str | None = None,
    ) -> tuple[str, list[tuple[Tag, int, int]]]:
        base_span = _html_span(element)
        table_id = _stable_id(
            "table", document.document_id, base_span.xpath, str(element)
        )
        context = cls._context_metadata(
            parent_table_id,
            parent_row_index,
            parent_column_index,
            parent_figure_id,
        )
        span = base_span.model_copy(
            update={"table_id": table_id, "metadata": dict(context)}
        )
        cells: list[CanonicalCell] = []
        cell_positions: list[tuple[Tag, int, int]] = []
        grid: list[list[str]] = []
        rowspan_until: dict[int, int] = {}
        direct_rows = [row for row in element.find_all("tr") if row.find_parent("table") is element]
        if len(direct_rows) > MAX_TABLE_ROWS:
            _raise_table_limit(document, "HTML", "row count")
        row_groups: list[Tag] = []
        for row in direct_rows:
            group = row.find_parent(["thead", "tbody", "tfoot"])
            if group is None or group.find_parent("table") is not element:
                group = element
            row_groups.append(group)
        group_last_row: dict[int, int] = {}
        for row_index, group in enumerate(row_groups):
            group_last_row[id(group)] = row_index
        header_row = False
        maximum_width = 0
        for row_index, row in enumerate(direct_rows):
            occupied = {
                column for column, last_row in rowspan_until.items() if last_row >= row_index
            }
            row_values: list[str] = [""] * (max(occupied, default=-1) + 1)
            column_index = 0
            direct_cells = [
                cell
                for cell in row.find_all(["th", "td"], recursive=False)
                if cell.find_parent("tr") is row
            ]
            if row_index == 0:
                header_row = any(cell.name == "th" for cell in direct_cells)
            for cell in direct_cells:
                remaining_group_rows = (
                    group_last_row[id(row_groups[row_index])] - row_index + 1
                )
                rowspan = cls._safe_span_value(
                    document,
                    cell,
                    "rowspan",
                    zero_value=remaining_group_rows,
                )
                rowspan = min(rowspan, remaining_group_rows)
                colspan = cls._safe_span_value(
                    document, cell, "colspan"
                )
                if colspan > MAX_TABLE_COLUMNS:
                    _raise_table_limit(document, "HTML", "column count")
                while any(
                    logical_column in occupied
                    for logical_column in range(column_index, column_index + colspan)
                ):
                    column_index += 1
                value = HtmlCanonicalAdapter._cell_text(element, cell)
                is_header = cell.name == "th"
                required = column_index + colspan
                if required > MAX_TABLE_COLUMNS:
                    _raise_table_limit(document, "HTML", "column count")
                maximum_width = max(maximum_width, required)
                if len(direct_rows) * maximum_width > MAX_TABLE_GRID_CELLS:
                    _raise_table_limit(document, "HTML", "grid cell count")
                if len(row_values) < required:
                    row_values.extend([""] * (required - len(row_values)))
                row_values[column_index] = value
                cells.append(
                    CanonicalCell(
                        text=value,
                        row_index=row_index,
                        column_index=column_index,
                        rowspan=rowspan,
                        colspan=colspan,
                        is_header=is_header,
                        source_spans=[
                            SourceSpan(
                                xpath=_html_xpath(cell),
                                css_selector=_html_selector(cell),
                                element_id=cell.get("id"),
                                table_id=table_id,
                                row_index=row_index,
                                column_index=column_index,
                                heading_path=builder.heading_path,
                                metadata=dict(context),
                            )
                        ],
                    )
                )
                cell_positions.append((cell, row_index, column_index))
                for occupied_column in range(column_index, column_index + colspan):
                    occupied.add(occupied_column)
                    if rowspan > 1:
                        rowspan_until[occupied_column] = row_index + rowspan - 1
                column_index += colspan
            grid.append(row_values)
        width = max((len(row) for row in grid), default=0)
        grid = [(row + [""] * width)[:width] for row in grid]
        headers = grid[0] if header_row and grid else []
        rows = grid[1:] if header_row else grid
        caption_element = element.find("caption")
        if caption_element is not None and caption_element.find_parent("table") is not element:
            caption_element = None
        caption = caption_element.get_text(" ", strip=True) if caption_element else None
        normalized = _table_markdown(headers, rows)
        table = CanonicalTable(
            table_id=table_id,
            caption=caption,
            headers=headers,
            rows=rows,
            cells=cells,
            source_html=str(element),
            normalized_markdown=normalized,
            source_spans=[span.model_copy(update={"heading_path": builder.heading_path})],
            metadata=dict(context),
        )
        document.tables.append(table)
        if caption_element is not None and caption:
            builder.add_content(
                caption,
                cls._context_span(caption_element, context),
                block_type="caption",
                table_id=table_id,
                metadata=dict(context),
            )
        builder.add_content(
            normalized,
            span,
            block_type="table",
            table_id=table_id,
            metadata=dict(context),
        )
        return table_id, cell_positions

    @staticmethod
    def _safe_span_value(
        document: CanonicalDocument,
        cell: Tag,
        attribute: str,
        *,
        zero_value: int | None = None,
    ) -> int:
        raw = cell.get(attribute, 1)
        try:
            value = int(raw)
        except (TypeError, ValueError):
            document.warnings.append(
                f"Invalid HTML {attribute}={raw!r}; using 1 at {_html_xpath(cell)}."
            )
            return 1
        if value == 0 and zero_value is not None:
            return max(1, zero_value)
        if value < 1:
            document.warnings.append(
                f"Invalid HTML {attribute}={raw!r}; using 1 at {_html_xpath(cell)}."
            )
            return 1
        return value

    @staticmethod
    def _cell_text(table: Tag, cell: Tag) -> str:
        pieces: list[str] = []

        def visit(node: Tag | NavigableString) -> None:
            if isinstance(node, NavigableString):
                if node.find_parent("table") is table and str(node).strip():
                    pieces.append(str(node).strip())
                return
            if not isinstance(node, Tag):
                return
            if node.name == "table":
                return
            if node.name == "img":
                value = node.get("alt") or node.get("src")
                if value:
                    pieces.append(value.strip())
                return
            if _is_math_element(node):
                annotation = node.find(
                    "annotation", attrs={"encoding": re.compile("tex", re.I)}
                )
                value = node.get("data-latex") or (
                    annotation.get_text("", strip=True)
                    if annotation
                    else node.get_text(" ", strip=True)
                )
                if value:
                    pieces.append(value.strip())
                return
            for child in list(node.children):
                visit(child)

        for child in list(cell.children):
            visit(child)
        return " ".join(pieces)

    def _add_figure(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        element: Tag,
        *,
        context: dict[str, object] | None = None,
    ) -> None:
        context = dict(context or {})
        span = self._context_span(element, context)
        image = next(
            (
                candidate
                for candidate in element.find_all("img")
                if candidate.find_parent("figure") is element
            ),
            None,
        )
        caption_element = next(
            (
                candidate
                for candidate in element.find_all("figcaption")
                if candidate.find_parent("figure") is element
            ),
            None,
        )
        caption = caption_element.get_text(" ", strip=True) if caption_element else None
        src = image.get("src") if image else None
        alt = image.get("alt") if image else None
        figure_span = span.model_copy(update={"heading_path": builder.heading_path})
        asset_path = _register_local_figure_asset(document, src, figure_span)
        figure_id = _stable_id("figure", document.document_id, span.xpath, src, caption)
        document.figures.append(
            CanonicalFigure(
                figure_id=figure_id,
                caption=caption or alt,
                asset_path=asset_path,
                source_spans=[figure_span],
                metadata={
                    "src": src,
                    "alt": alt,
                    "source_html": str(element),
                    **context,
                },
            )
        )
        child_context = {**context, "parent_figure_id": figure_id}
        emitted_figure = False
        narrative_index = 0

        def emit_figure() -> None:
            nonlocal emitted_figure
            if emitted_figure:
                return
            builder.add_content(
                alt or src or "Figure",
                span,
                block_type="figure",
                figure_id=figure_id,
                metadata=dict(context),
                retrievable=(
                    False if context.get("retrieval_covered_by_table_id") else None
                ),
            )
            emitted_figure = True

        if image is None:
            emit_figure()

        def visit(node: Tag | NavigableString) -> None:
            nonlocal narrative_index
            if isinstance(node, NavigableString):
                text = str(node)
                if text.strip():
                    parent = node.parent if isinstance(node.parent, Tag) else element
                    text_span = self._context_span(parent, child_context)
                    builder.add_content(
                        text,
                        text_span.model_copy(
                            update={
                                "metadata": {
                                    **child_context,
                                    "segment_index": narrative_index,
                                }
                            }
                        ),
                        metadata=dict(child_context),
                        retrievable=(
                            False
                            if child_context.get("retrieval_covered_by_table_id")
                            else None
                        ),
                    )
                    narrative_index += 1
                return
            if not isinstance(node, Tag):
                return
            if node is image:
                emit_figure()
                return
            if node is caption_element:
                if caption:
                    builder.add_content(
                        caption,
                        self._context_span(node, context),
                        block_type="caption",
                        figure_id=figure_id,
                        metadata=dict(context),
                        retrievable=(
                            False
                            if context.get("retrieval_covered_by_table_id")
                            else None
                        ),
                    )
                return
            if node.name == "p":
                self._walk_inline_container(
                    document,
                    builder,
                    node,
                    context=child_context,
                )
                return
            if node.name == "pre":
                builder.add_content(
                    node.get_text("", strip=False).strip("\r\n"),
                    self._context_span(node, child_context),
                    metadata={
                        "kind": "code",
                        "language": self._code_language(node.find("code")),
                        **child_context,
                    },
                    retrievable=(
                        False
                        if child_context.get("retrieval_covered_by_table_id")
                        else None
                    ),
                )
                return
            if node.name == "code":
                builder.add_content(
                    node.get_text("", strip=False),
                    self._context_span(node, child_context),
                    metadata={
                        "kind": "code",
                        "language": self._code_language(node),
                        **child_context,
                    },
                    retrievable=(
                        False
                        if child_context.get("retrieval_covered_by_table_id")
                        else None
                    ),
                )
                return
            if _is_math_element(node):
                self._add_formula(document, builder, node, context=child_context)
                return
            if node.name == "img":
                self._add_image(document, builder, node, context=child_context)
                return
            if node.name == "table":
                self._walk_table_tree(
                    document,
                    builder,
                    node,
                    parent_table_id=context.get("parent_table_id"),
                    parent_row_index=context.get("parent_row_index"),
                    parent_column_index=context.get("parent_column_index"),
                    parent_figure_id=figure_id,
                )
                return
            if node.name == "figure":
                self._add_figure(document, builder, node, context=child_context)
                return
            for child in list(node.children):
                visit(child)

        for child in list(element.children):
            visit(child)
        emit_figure()

    @classmethod
    def _add_image(
        cls,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        element: Tag,
        *,
        context: dict[str, object] | None = None,
    ) -> None:
        context = dict(context or {})
        span = cls._context_span(element, context)
        src = element.get("src")
        alt = element.get("alt")
        figure_span = span.model_copy(update={"heading_path": builder.heading_path})
        asset_path = _register_local_figure_asset(document, src, figure_span)
        figure_id = _stable_id("figure", document.document_id, span.xpath, src)
        document.figures.append(
            CanonicalFigure(
                figure_id=figure_id,
                caption=alt,
                asset_path=asset_path,
                source_spans=[figure_span],
                metadata={
                    "src": src,
                    "alt": alt,
                    "source_html": str(element),
                    **context,
                },
            )
        )
        builder.add_content(
            alt or src or "Image",
            span,
            block_type="figure",
            figure_id=figure_id,
            metadata=dict(context),
            retrievable=(
                False if context.get("retrieval_covered_by_table_id") else None
            ),
        )

    @classmethod
    def _add_formula(
        cls,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        element: Tag,
        *,
        context: dict[str, object] | None = None,
    ) -> None:
        context = dict(context or {})
        span = cls._context_span(element, context)
        annotation = element.find("annotation", attrs={"encoding": re.compile("tex", re.I)})
        latex = element.get("data-latex") or (
            annotation.get_text("", strip=True) if annotation else element.get_text(" ", strip=True)
        )
        formula_id = _stable_id("formula", document.document_id, span.xpath, latex)
        document.formulas.append(
            CanonicalFormula(
                formula_id=formula_id,
                latex=latex,
                source_spans=[span.model_copy(update={"heading_path": builder.heading_path})],
                metadata={
                    "source_format": "html",
                    "source_html": str(element),
                    **context,
                },
            )
        )
        builder.add_content(
            latex,
            span,
            block_type="formula",
            formula_id=formula_id,
            metadata=dict(context),
            retrievable=(
                False if context.get("retrieval_covered_by_table_id") else None
            ),
        )


@dataclass(frozen=True)
class _DocxEvent:
    kind: str
    text: str = ""
    element: object | None = None
    relationship_id: str | None = None
    metadata: dict[str, str] | None = None


class DocxCanonicalAdapter:
    parser_source = "docx"

    def __init__(self, asset_cache_root: Path | None = None) -> None:
        self.asset_cache_root = Path(asset_cache_root) if asset_cache_root is not None else None

    def parse(self, path: Path) -> CanonicalDocument:
        path = _validate_path(path)
        try:
            docx = DocxDocument(str(path))
        except Exception as exc:
            raise _parse_error(path, f"Unable to read DOCX: {exc}", exc)
        document = _new_document(
            path,
            self.parser_source,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        document.title = docx.core_properties.title or display_title_from_path(path)
        builder = _DocumentBuilder(document)
        paragraph_index = 0
        table_index = 0

        for child, locator in self._iter_docx_blocks(docx.element.body, "body", path):
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, docx)
                paragraph_id = child.get(qn("w14:paraId")) or f"paragraph-{paragraph_index}"
                self._emit_paragraph(
                    document,
                    builder,
                    docx,
                    paragraph,
                    SourceSpan(paragraph_id=paragraph_id),
                )
                paragraph_index += 1
            elif isinstance(child, CT_Tbl):
                table_count = self._add_table(
                    document,
                    builder,
                    docx,
                    Table(child, docx),
                    locator=locator,
                    nesting_depth=1,
                )
                table_index += table_count

        document.parser_metadata.update(
            {
                "paragraph_count": paragraph_index,
                "paragraphs": paragraph_index,
                "table_count": table_index,
            }
        )
        if not document.blocks:
            document.warnings.append("DOCX contains no supported content blocks.")
        return document

    @classmethod
    def _iter_docx_blocks(cls, container, prefix: str, source_path: Path):
        stack: list[tuple[object, str, int]] = []
        children = list(container.iterchildren())
        for child_index in reversed(range(len(children))):
            stack.append((children[child_index], f"{prefix}/{child_index}", 1))
        while stack:
            child, locator, depth = stack.pop()
            if depth > MAX_NESTING_DEPTH:
                raise _parse_error(
                    source_path,
                    f"DOCX nesting depth exceeds the {MAX_NESTING_DEPTH} element limit.",
                )
            if isinstance(child, (CT_P, CT_Tbl)):
                yield child, locator
                continue
            descendants = list(child.iterchildren())
            for child_index in reversed(range(len(descendants))):
                stack.append(
                    (
                        descendants[child_index],
                        f"{locator}/{child_index}",
                        depth + 1,
                    )
                )

    def _emit_paragraph(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        docx,
        paragraph: Paragraph,
        span: SourceSpan,
        *,
        structures_only: bool = False,
    ) -> None:
        events = list(
            self._paragraph_events(paragraph, Path(document.source_path or "document.docx"))
        )
        style_name = paragraph.style.name if paragraph.style is not None else ""
        heading_match = re.match(r"Heading\s+(\d+)", style_name, re.I)
        heading_text = "".join(self._heading_event_text(event) for event in events)
        if not structures_only and heading_match and heading_text.strip():
            heading_block = builder.add_heading(
                heading_text, int(heading_match.group(1)), span, style=style_name
            )
            for event_index, event in enumerate(events):
                if event.kind != "text":
                    self._emit_structure_event(
                        document,
                        builder,
                        docx,
                        event,
                        span,
                        event_index,
                        parent_heading_block_id=heading_block.block_id,
                        inline_event_index=event_index,
                    )
            return
        if not structures_only and style_name.lower() == "title" and heading_text.strip():
            document.title = heading_text.strip()
            heading_block = builder.add_heading(heading_text, 1, span, style=style_name)
            for event_index, event in enumerate(events):
                if event.kind != "text":
                    self._emit_structure_event(
                        document,
                        builder,
                        docx,
                        event,
                        span,
                        event_index,
                        parent_heading_block_id=heading_block.block_id,
                        inline_event_index=event_index,
                    )
            return

        text_buffer: list[str] = []

        def flush() -> None:
            text = "".join(text_buffer)
            text_buffer.clear()
            if not structures_only and text.strip():
                builder.add_content(text, span, metadata={"style": style_name})

        for event_index, event in enumerate(events):
            if event.kind == "text":
                text_buffer.append(event.text)
                continue
            flush()
            self._emit_structure_event(document, builder, docx, event, span, event_index)
        flush()

    @staticmethod
    def _heading_event_text(event: _DocxEvent) -> str:
        if event.kind == "text":
            return event.text
        if event.kind == "formula":
            return event.text.strip() or "[OMML formula]"
        if event.kind == "image":
            metadata = event.metadata or {}
            return (
                metadata.get("descr")
                or metadata.get("title")
                or metadata.get("name")
                or "[Image]"
            )
        return ""

    @classmethod
    def _paragraph_events(cls, paragraph: Paragraph, source_path: Path):
        validation_stack = [(paragraph._p, 0)]
        while validation_stack:
            node, depth = validation_stack.pop()
            if depth > MAX_NESTING_DEPTH:
                raise _parse_error(
                    source_path,
                    f"DOCX nesting depth exceeds the {MAX_NESTING_DEPTH} element limit.",
                )
            validation_stack.extend(
                (child, depth + 1) for child in reversed(list(node.iterchildren()))
            )

        stack = list(reversed(list(paragraph._p.iterchildren())))
        while stack:
            node = stack.pop()
            if node.tag in {qn("m:oMath"), qn("m:oMathPara")}:
                omml = etree.tostring(node, encoding="unicode")
                text = "".join(item.text or "" for item in node.iter(qn("m:t"))).strip()
                yield _DocxEvent("formula", text=text, element=node, metadata={"omml": omml})
                continue
            if node.tag in {qn("w:drawing"), qn("w:pict"), qn("w:object")}:
                doc_properties = next(iter(node.iter(qn("wp:docPr"))), None)
                metadata = {
                    key: doc_properties.get(key)
                    for key in ("descr", "title", "name")
                    if doc_properties is not None and doc_properties.get(key)
                }
                for blip in node.iter(qn("a:blip")):
                    embedded_id = blip.get(qn("r:embed"))
                    linked_id = blip.get(qn("r:link"))
                    yield _DocxEvent(
                        "image",
                        element=blip,
                        relationship_id=embedded_id or linked_id,
                        metadata={
                            **metadata,
                            "relationship_attribute": "embed" if embedded_id else "link",
                        },
                    )
                continue
            if node.tag in {qn("w:t"), qn("w:instrText")}:
                yield _DocxEvent("text", text=node.text or "")
                continue
            if node.tag == qn("w:tab"):
                yield _DocxEvent("text", text="\t")
                continue
            if node.tag in {qn("w:br"), qn("w:cr")}:
                yield _DocxEvent("text", text="\n")
                continue
            stack.extend(reversed(list(node.iterchildren())))

    def _emit_structure_event(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        docx,
        event: _DocxEvent,
        span: SourceSpan,
        event_index: int,
        *,
        parent_heading_block_id: str | None = None,
        inline_event_index: int | None = None,
    ) -> None:
        if event.kind == "image":
            self._add_image(
                document,
                builder,
                docx,
                event,
                span,
                event_index,
                parent_heading_block_id=parent_heading_block_id,
                inline_event_index=inline_event_index,
            )
        elif event.kind == "formula":
            self._add_formula(
                document,
                builder,
                event,
                span,
                event_index,
                parent_heading_block_id=parent_heading_block_id,
                inline_event_index=inline_event_index,
            )

    def _add_image(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        docx,
        event: _DocxEvent,
        span: SourceSpan,
        event_index: int,
        *,
        parent_heading_block_id: str | None = None,
        inline_event_index: int | None = None,
    ) -> None:
        relationship_id = event.relationship_id
        if not relationship_id or relationship_id not in docx.part.rels:
            document.warnings.append(
                f"DOCX image relationship {relationship_id or '<missing>'!r} is missing; image skipped."
            )
            return
        relationship = docx.part.rels[relationship_id]
        try:
            source_target = str(relationship.target_ref).replace("\\", "/")
        except Exception as exc:
            raise _parse_error(
                Path(document.source_path or "document.docx"),
                f"Unable to read DOCX image relationship {relationship_id}: {exc}",
                exc,
            )
        metadata = event.metadata or {}
        caption = metadata.get("descr") or metadata.get("title") or metadata.get("name")
        heading_link = {
            key: value
            for key, value in {
                "parent_heading_block_id": parent_heading_block_id,
                "inline_event_index": inline_event_index,
            }.items()
            if value is not None
        }
        image_span = span.model_copy(
            update={
                "image_relationship_id": relationship_id,
                "heading_path": builder.heading_path,
                "metadata": {
                    **span.metadata,
                    "relationship_target": source_target,
                    **heading_link,
                },
            }
        )
        figure_id = _stable_id(
            "figure",
            document.document_id,
            span.paragraph_id,
            span.table_id,
            span.row_index,
            span.column_index,
            relationship_id,
            event_index,
        )
        if relationship.is_external:
            document.warnings.append(
                f"DOCX image relationship {relationship_id!r} is external and was not materialized."
            )
            document.figures.append(
                CanonicalFigure(
                    figure_id=figure_id,
                    caption=caption,
                    asset_path=None,
                    source_spans=[image_span],
                    metadata={
                        "relationship_id": relationship_id,
                        "relationship_target": source_target,
                        "external": True,
                        **metadata,
                        **heading_link,
                    },
                )
            )
            builder.add_content(
                caption or source_target or "External image",
                image_span,
                block_type="figure",
                figure_id=figure_id,
                metadata={
                    "relationship_id": relationship_id,
                    "relationship_target": source_target,
                    "external": True,
                    **metadata,
                    **heading_link,
                },
            )
            return

        try:
            target_part = relationship.target_part
            blob = target_part.blob
            asset_sha = hashlib.sha256(blob).hexdigest()
            target_name = Path(str(target_part.partname)).name
            safe_name = self._safe_asset_name(target_name, asset_sha)
            asset_path = f"assets/{safe_name}"
            source_path = self._materialize_asset(
                document, safe_name, blob, asset_sha
            )
            media_type = getattr(target_part, "content_type", None) or (
                mimetypes.guess_type(target_name)[0] or "application/octet-stream"
            )
        except Exception as exc:
            from app.services.parser import DocumentParseError

            if isinstance(exc, DocumentParseError):
                raise
            raise _parse_error(
                Path(document.source_path or "document.docx"),
                f"Unable to read embedded DOCX image {relationship_id}: {exc}",
                exc,
            )

        asset_id = _stable_id("asset", document.document_id, asset_sha)
        if not any(asset.asset_id == asset_id for asset in document.assets):
            document.assets.append(
                CanonicalAsset(
                    asset_id=asset_id,
                    path=asset_path,
                    media_type=media_type,
                    sha256=asset_sha,
                    source_path=str(source_path),
                    source_spans=[image_span],
                    metadata={
                        "relationship_id": relationship_id,
                        "relationship_target": source_target,
                        "embedded": True,
                        "size_bytes": len(blob),
                        "cache_limit_bytes": ADAPTER_ASSET_CACHE_MAX_BYTES,
                        "cache_cleanup_policy": ADAPTER_ASSET_CACHE_CLEANUP_POLICY,
                    },
                )
            )
        document.figures.append(
            CanonicalFigure(
                figure_id=figure_id,
                caption=caption,
                asset_path=asset_path,
                source_spans=[image_span],
                metadata={
                    "asset_id": asset_id,
                    "relationship_id": relationship_id,
                    "relationship_target": source_target,
                    **metadata,
                    **heading_link,
                },
            )
        )
        builder.add_content(
            caption or target_name,
            image_span,
            block_type="figure",
            figure_id=figure_id,
            metadata={"relationship_id": relationship_id, **metadata, **heading_link},
        )

    def _materialize_asset(
        self,
        document: CanonicalDocument,
        safe_name: str,
        blob: bytes,
        expected_sha: str,
    ) -> Path:
        return _materialize_adapter_asset(
            document,
            safe_name,
            blob,
            expected_sha,
            source_kind="DOCX",
            asset_cache_root=self.asset_cache_root,
        )

    @staticmethod
    def _safe_asset_name(target_name: str, asset_sha: str) -> str:
        suffix = Path(target_name).suffix.lower()
        if not re.fullmatch(r"\.[a-z0-9]{1,10}", suffix):
            suffix = ""
        return f"image-{asset_sha[:24]}{suffix}"

    @staticmethod
    def _add_formula(
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        event: _DocxEvent,
        span: SourceSpan,
        event_index: int,
        *,
        parent_heading_block_id: str | None = None,
        inline_event_index: int | None = None,
    ) -> None:
        omml = (event.metadata or {}).get("omml", "")
        source_text = event.text.strip()
        heading_link = {
            key: value
            for key, value in {
                "parent_heading_block_id": parent_heading_block_id,
                "inline_event_index": inline_event_index,
            }.items()
            if value is not None
        }
        formula_id = _stable_id(
            "formula",
            document.document_id,
            span.paragraph_id,
            span.table_id,
            span.row_index,
            span.column_index,
            event_index,
            omml,
        )
        formula_span = span.model_copy(
            update={
                "heading_path": builder.heading_path,
                "metadata": {
                    **span.metadata,
                    "formula_event_index": event_index,
                    **heading_link,
                },
            }
        )
        document.formulas.append(
            CanonicalFormula(
                formula_id=formula_id,
                latex=source_text or "[OMML formula]",
                source_spans=[formula_span],
                warnings=["OMML source text is preserved but is not normalized LaTeX."],
                metadata={"source_format": "omml", "omml": omml, **heading_link},
            )
        )
        builder.add_content(
            source_text or "[OMML formula]",
            formula_span,
            block_type="formula",
            formula_id=formula_id,
            metadata={"source_format": "omml", **heading_link},
        )

    def _add_table(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        docx,
        table: Table,
        *,
        locator: str,
        parent_table_id: str | None = None,
        parent_row_index: int | None = None,
        parent_column_index: int | None = None,
        nesting_depth: int = 1,
    ) -> int:
        table_id = _stable_id("table", document.document_id, "docx", locator)
        context = {
            key: value
            for key, value in {
                "locator": locator,
                "parent_table_id": parent_table_id,
                "parent_row_index": parent_row_index,
                "parent_column_index": parent_column_index,
            }.items()
            if value is not None
        }
        cells: list[CanonicalCell] = []
        grid: list[list[str]] = []
        active_merges: dict[int, CanonicalCell] = {}
        cell_entries: list[tuple[_Cell, int, int, str]] = []
        raw_rows = list(table._tbl.findall(qn("w:tr")))
        if len(raw_rows) > MAX_TABLE_ROWS:
            _raise_table_limit(document, "DOCX", "row count")

        maximum_width = 0
        for row_index, raw_row in enumerate(raw_rows):
            row_values: list[str] = []
            column_index = self._grid_before(raw_row)
            if column_index > MAX_TABLE_COLUMNS:
                _raise_table_limit(document, "DOCX", "column count")
            maximum_width = max(maximum_width, column_index)
            if len(raw_rows) * maximum_width > MAX_TABLE_GRID_CELLS:
                _raise_table_limit(document, "DOCX", "grid cell count")
            if column_index:
                row_values.extend([""] * column_index)
            continued: set[int] = set()
            restarted: set[int] = set()
            for raw_cell in raw_row.findall(qn("w:tc")):
                cell = _Cell(raw_cell, table)
                colspan = self._docx_grid_span(raw_cell)
                if colspan > MAX_TABLE_COLUMNS:
                    _raise_table_limit(document, "DOCX", "column count")
                vmerge = raw_cell.tcPr.vMerge
                merge_value = str(vmerge.val).lower() if vmerge is not None else ""
                is_restart = vmerge is not None and merge_value == "restart"
                is_continue = vmerge is not None and not is_restart
                required = column_index + colspan
                if required > MAX_TABLE_COLUMNS:
                    _raise_table_limit(document, "DOCX", "column count")
                maximum_width = max(maximum_width, required)
                if len(raw_rows) * maximum_width > MAX_TABLE_GRID_CELLS:
                    _raise_table_limit(document, "DOCX", "grid cell count")
                if len(row_values) < required:
                    row_values.extend([""] * (required - len(row_values)))
                cell_locator = f"{locator}/cell-{row_index}-{column_index}"
                cell_entries.append(
                    (cell, row_index, column_index, cell_locator)
                )

                if is_continue and column_index in active_merges:
                    origin = active_merges[column_index]
                    origin.rowspan += 1
                    for logical_column in range(column_index, column_index + colspan):
                        continued.add(logical_column)
                    column_index += colspan
                    continue

                text = cell.text
                row_values[column_index] = text
                cell_span = SourceSpan(
                    table_id=table_id,
                    row_index=row_index,
                    column_index=column_index,
                    heading_path=builder.heading_path,
                    metadata={"locator": cell_locator, **context},
                )
                canonical_cell = CanonicalCell(
                    text=text,
                    row_index=row_index,
                    column_index=column_index,
                    colspan=colspan,
                    is_header=row_index == 0,
                    source_spans=[cell_span],
                )
                cells.append(canonical_cell)
                for logical_column in range(column_index, column_index + colspan):
                    active_merges.pop(logical_column, None)
                    if is_restart:
                        active_merges[logical_column] = canonical_cell
                        restarted.add(logical_column)
                column_index += colspan
            active_merges = {
                column: origin
                for column, origin in active_merges.items()
                if column in continued or column in restarted
            }
            grid.append(row_values)

        width = max((len(row) for row in grid), default=0)
        grid = [(row + [""] * width)[:width] for row in grid]
        headers = grid[0] if grid else []
        rows = grid[1:] if grid else []
        span = SourceSpan(table_id=table_id, metadata=dict(context))
        normalized = _table_markdown(headers, rows)
        document.tables.append(
            CanonicalTable(
                table_id=table_id,
                headers=headers,
                rows=rows,
                cells=cells,
                normalized_markdown=normalized,
                source_spans=[span.model_copy(update={"heading_path": builder.heading_path})],
                metadata=dict(context),
            )
        )
        builder.add_content(
            normalized,
            span,
            block_type="table",
            table_id=table_id,
            metadata=dict(context),
        )

        nested_table_count = 0
        for cell, row_index, column_index, cell_locator in cell_entries:
            nested_table_count += self._walk_docx_cell_contents(
                document,
                builder,
                docx,
                cell,
                prefix=cell_locator,
                parent_table_id=table_id,
                parent_row_index=row_index,
                parent_column_index=column_index,
                base_depth=nesting_depth,
            )
        return 1 + nested_table_count

    def _walk_docx_cell_contents(
        self,
        document: CanonicalDocument,
        builder: _DocumentBuilder,
        docx,
        cell: _Cell,
        *,
        prefix: str,
        parent_table_id: str,
        parent_row_index: int,
        parent_column_index: int,
        base_depth: int,
    ) -> int:
        nested_table_count = 0

        stack: list[tuple[object, str, int]] = []
        children = list(cell._tc.iterchildren())
        for child_index in reversed(range(len(children))):
            stack.append(
                (children[child_index], f"{prefix}/{child_index}", base_depth + 1)
            )
        while stack:
            child, child_locator, depth = stack.pop()
            if depth > MAX_NESTING_DEPTH:
                raise _parse_error(
                    Path(document.source_path or "document.docx"),
                    f"DOCX nesting depth exceeds the {MAX_NESTING_DEPTH} element limit.",
                )
            if isinstance(child, CT_P):
                paragraph = Paragraph(child, cell)
                paragraph_id = child.get(qn("w14:paraId")) or child_locator.replace(
                    "/", "-"
                )
                self._emit_paragraph(
                    document,
                    builder,
                    docx,
                    paragraph,
                    SourceSpan(
                        paragraph_id=paragraph_id,
                        table_id=parent_table_id,
                        row_index=parent_row_index,
                        column_index=parent_column_index,
                        metadata={"locator": child_locator},
                    ),
                    structures_only=True,
                )
                continue
            if isinstance(child, CT_Tbl):
                nested_table_count += self._add_table(
                    document,
                    builder,
                    docx,
                    Table(child, cell),
                    locator=child_locator,
                    parent_table_id=parent_table_id,
                    parent_row_index=parent_row_index,
                    parent_column_index=parent_column_index,
                    nesting_depth=depth,
                )
                continue
            descendants = list(child.iterchildren())
            for descendant_index in reversed(range(len(descendants))):
                stack.append(
                    (
                        descendants[descendant_index],
                        f"{child_locator}/{descendant_index}",
                        depth + 1,
                    )
                )
        return nested_table_count

    @staticmethod
    def _docx_grid_span(raw_cell) -> int:
        grid_span = raw_cell.tcPr.gridSpan
        if grid_span is None:
            return 1
        try:
            return max(1, int(grid_span.val))
        except (TypeError, ValueError):
            return 1

    @staticmethod
    def _grid_before(raw_row) -> int:
        grid_before = raw_row.find(f"./{qn('w:trPr')}/{qn('w:gridBefore')}")
        if grid_before is None:
            return 0
        try:
            return max(0, int(grid_before.get(qn("w:val"), "0")))
        except ValueError:
            return 0


def run_mineru(path: Path, page_count: int) -> CanonicalDocument | None:
    """Run MinerU without entering the canonical PDF dispatcher recursively."""
    from app.services import parser

    parsed = parser._parse_pdf_with_mineru(path, page_count)
    if parsed is None:
        return None
    return _parsed_pdf_to_canonical(path, parsed, "mineru", page_count)


def run_document_intelligence(
    path: Path,
    page_count: int,
    page_texts: list[str],
    page_indices: set[int] | list[int] | tuple[int, ...] | None = None,
) -> CanonicalDocument | None:
    """Run full or page-scoped Document Intelligence and canonicalize its output."""
    from app.services import parser

    parsed = parser._parse_pdf_with_document_intelligence(
        path,
        page_texts,
        page_count,
        page_indices=page_indices,
    )
    if parsed is None:
        return None
    document = _parsed_pdf_to_canonical(
        path,
        parsed,
        "document_intelligence",
        page_count,
    )
    intelligence = parsed.metadata.get("document_intelligence", {})
    page_outputs = intelligence.get("page_outputs", []) if isinstance(intelligence, dict) else []
    candidates: list[str] = []
    if isinstance(page_outputs, list):
        for output in page_outputs:
            if not isinstance(output, dict):
                continue
            sections = output.get("sections", [])
            if isinstance(sections, list):
                candidates.extend(value for value in sections if isinstance(value, str))
    selected = set(page_indices) if page_indices is not None else {0, 1}
    candidates.extend(
        text
        for index, text in enumerate(page_texts)
        if index in selected and index < 2 and isinstance(text, str)
    )
    document.abstract = _extract_explicit_abstract(candidates)
    return document


def run_text_layer_fallback(
    path: Path,
    page_count: int,
    page_texts: list[str] | None = None,
    *,
    text_layer_warnings: list[str] | None = None,
) -> CanonicalDocument:
    """Build the final canonical fallback from complete pypdf page text."""
    from app.services import parser

    warnings = list(text_layer_warnings or [])
    if page_texts is None:
        try:
            page_texts, extracted_count = parser._extract_pdf_text_layer(path)
            page_count = max(page_count, extracted_count)
        except Exception as exc:  # noqa: BLE001
            page_texts = [""] * page_count
            warnings.append(f"Text layer extraction failed: {exc}")
    document = _new_document(path, "pypdf_text_layer", "application/pdf")
    document.parser_source = "pypdf_text_layer"
    document.parser_metadata.update(
        {"adapter": "pypdf_text_layer", "format": "pdf", "parser_mode": "pdf_text_layer"}
    )
    builder = _DocumentBuilder(document)
    for page_index in range(page_count):
        text = page_texts[page_index] if page_index < len(page_texts) else ""
        if not text.strip():
            continue
        builder.add_content(
            text,
            SourceSpan(
                page_index=page_index,
                page_label=str(page_index + 1),
                source_block_id=f"pypdf-page-{page_index + 1}",
            ),
            metadata={"source": "pypdf_text_layer"},
        )
    document.metadata.update(
        {
            "expected_page_count": page_count,
            "parsed_page_indices": list(range(page_count)),
            "text_layer_pages": list(page_texts),
            "text_layer_warnings": warnings,
        }
    )
    document.warnings.extend(warnings)
    return document


def _parsed_pdf_to_canonical(
    path: Path,
    parsed,
    parser_source: str,
    page_count: int,
) -> CanonicalDocument:
    document = _new_document(path, parser_source, "application/pdf")
    document.parser_source = parser_source
    document.title = parsed.title
    document.metadata.update(
        {
            "expected_page_count": page_count,
            "legacy_pdf_metadata": parsed.metadata,
        }
    )
    document.parser_metadata.update(
        {
            "adapter": parser_source,
            "format": "pdf",
            "parser_mode": parsed.metadata.get("parser_mode", parser_source),
            "source_parser_metadata": parsed.metadata,
        }
    )
    intelligence = parsed.metadata.get("document_intelligence", {})
    if not isinstance(intelligence, dict):
        intelligence = {}

    table_entries: list[tuple[str | None, str, CanonicalTable]] = []
    for index, raw_table in enumerate(intelligence.get("tables", [])):
        if not isinstance(raw_table, dict):
            continue
        page_label = _clean_page_label(raw_table.get("page_label"))
        table_text = str(raw_table.get("markdown") or raw_table.get("text") or "").strip()
        caption, headers, rows = _parse_pdf_table_markdown(table_text)
        span = _pdf_span(page_label, f"{parser_source}-table-{index + 1}")
        table_id = _stable_id(
            "table", document.document_id, page_label, index, table_text
        )
        table = CanonicalTable(
            table_id=table_id,
            caption=caption,
            headers=headers,
            rows=rows,
            cells=[
                CanonicalCell(
                    text=value,
                    row_index=row_index,
                    column_index=column_index,
                    is_header=row_index == 0,
                    source_spans=[span],
                )
                for row_index, row in enumerate([headers, *rows])
                for column_index, value in enumerate(row)
            ],
            source_markdown=table_text or None,
            normalized_markdown=_table_markdown(headers, rows) or table_text or None,
            source_spans=[span],
            status="accepted_mineru" if parser_source == "mineru" else "repaired_by_vision",
            metadata={
                key: value
                for key, value in raw_table.items()
                if key not in {"markdown", "text"}
            },
        )
        document.tables.append(table)
        table_entries.append((page_label, table_text, table))

    formula_entries: list[tuple[str | None, str, CanonicalFormula]] = []
    for index, raw_formula in enumerate(intelligence.get("formulas", [])):
        if not isinstance(raw_formula, dict):
            continue
        page_label = _clean_page_label(raw_formula.get("page_label"))
        formula_text = str(raw_formula.get("text") or raw_formula.get("latex") or "").strip()
        if not formula_text:
            continue
        caption = str(raw_formula.get("caption") or "").strip() or None
        description = str(
            raw_formula.get("description") or raw_formula.get("note") or ""
        ).strip() or None
        span = _pdf_span(page_label, f"{parser_source}-formula-{index + 1}")
        formula = CanonicalFormula(
            formula_id=_stable_id(
                "formula", document.document_id, page_label, index, formula_text
            ),
            latex=formula_text,
            caption=caption,
            description=description,
            source_spans=[span],
            metadata={
                key: value
                for key, value in raw_formula.items()
                if key not in {"text", "latex", "caption", "description", "note"}
            },
        )
        document.formulas.append(formula)
        formula_entries.append((page_label, formula_text, formula))

    figure_entries: list[tuple[str | None, str, CanonicalFigure]] = []
    for index, raw_figure in enumerate(intelligence.get("figures", [])):
        if not isinstance(raw_figure, dict):
            continue
        page_label = _clean_page_label(raw_figure.get("page_label"))
        caption = str(raw_figure.get("caption") or "").strip() or None
        description = str(raw_figure.get("note") or raw_figure.get("description") or "").strip() or None
        span = _pdf_span(page_label, f"{parser_source}-figure-{index + 1}")
        asset_path = _register_pdf_figure_asset(
            document,
            raw_figure,
            intelligence,
            span,
        )
        figure = CanonicalFigure(
            figure_id=_stable_id(
                "figure", document.document_id, page_label, index, caption, description
            ),
            caption=caption,
            description=description,
            asset_path=asset_path,
            source_spans=[span],
            metadata=dict(raw_figure),
        )
        document.figures.append(figure)
        figure_text = _pdf_figure_text(raw_figure)
        figure_entries.append((page_label, figure_text, figure))

    builder = _DocumentBuilder(document)
    chunks = list(parsed.chunks)
    if not chunks and parsed.text:
        from app.services.parser import ParsedChunk

        chunks = [ParsedChunk(ordinal=0, text=parsed.text)]

    used_tables: set[str] = set()
    used_formulas: set[str] = set()
    used_figures: set[str] = set()
    for chunk in chunks:
        text = str(chunk.text)
        if not text.strip():
            continue
        page_label = _clean_page_label(chunk.page_label)
        span = _pdf_span(
            page_label,
            str(chunk.heading or f"{parser_source}-chunk-{chunk.ordinal}"),
        )
        heading = str(chunk.heading or "").lower()
        table = _match_pdf_structure(
            table_entries,
            page_label,
            text,
            used_tables,
            force="table" in heading,
        )
        formula = _match_pdf_structure(
            formula_entries,
            page_label,
            text,
            used_formulas,
            force="formula" in heading or "equation" in heading,
        )
        figure = _match_pdf_structure(
            figure_entries,
            page_label,
            text,
            used_figures,
            force="figure" in heading or "image" in heading,
        )
        metadata = {
            "source_chunk_ordinal": chunk.ordinal,
            "source_heading": chunk.heading,
        }
        if table is not None:
            builder.add_content(
                text,
                span,
                block_type="table",
                table_id=table.table_id,
                metadata=metadata,
            )
        elif formula is not None:
            builder.add_content(
                text,
                span,
                block_type="formula",
                formula_id=formula.formula_id,
                metadata=metadata,
            )
        elif figure is not None:
            builder.add_content(
                text,
                span,
                block_type="figure",
                figure_id=figure.figure_id,
                metadata=metadata,
            )
        elif _is_standalone_pdf_heading(text):
            title = text.lstrip("#").strip()
            builder.add_heading(title, min(6, max(1, len(text) - len(text.lstrip("#")))), span, **metadata)
        else:
            block = builder.add_content(text, span, metadata=metadata)
            if chunk.heading:
                block.section_path = [str(chunk.heading)]
                block.source_spans[0].heading_path = [str(chunk.heading)]

    _append_unmatched_structures(
        document,
        builder,
        table_entries,
        formula_entries,
        figure_entries,
        used_tables,
        used_formulas,
        used_figures,
    )
    parsed_pages = sorted(
        {
            span.page_index
            for block in document.blocks
            for span in block.source_spans
            if span.page_index is not None
        }
    )
    document.metadata["parsed_page_indices"] = parsed_pages
    explicit_candidates = [chunk.text for chunk in chunks]
    document.abstract = _extract_explicit_abstract(explicit_candidates)
    if not document.blocks:
        document.warnings.append(f"{parser_source} returned no canonical content blocks.")
    return document


def _parse_pdf_table_markdown(text: str) -> tuple[str | None, list[str], list[list[str]]]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for header_index in range(max(0, len(lines) - 1)):
        if "|" not in lines[header_index] or not _is_table_separator(lines[header_index + 1]):
            continue
        headers = _split_pipe_row(lines[header_index])
        rows: list[list[str]] = []
        for line in lines[header_index + 2 :]:
            if "|" not in line:
                break
            rows.append(_split_pipe_row(line))
        caption_text = " ".join(lines[:header_index]).strip()
        return caption_text or None, headers, rows
    return None, [], []


def _clean_page_label(value: object) -> str | None:
    label = str(value or "").strip()
    return label if label and label != "?" else None


def _pdf_span(page_label: str | None, source_block_id: str) -> SourceSpan:
    page_index = int(page_label) - 1 if page_label and page_label.isdigit() else None
    return SourceSpan(
        page_index=page_index,
        page_label=page_label,
        source_block_id=source_block_id,
    )


def _pdf_figure_text(figure: dict) -> str:
    parts = [
        str(figure.get(key) or "").strip()
        for key in ("caption", "note", "description", "image_path", "path")
    ]
    return "\n".join(part for part in parts if part)


def _register_pdf_figure_asset(
    document: CanonicalDocument,
    figure: dict,
    intelligence: dict,
    span: SourceSpan,
) -> str | None:
    raw_path = str(figure.get("image_path") or figure.get("path") or "").strip()
    relative = Path(raw_path)
    if (
        not raw_path
        or relative.is_absolute()
        or ".." in relative.parts
        or "://" in raw_path
    ):
        return None

    roots: list[Path] = []
    content_list_path = intelligence.get("content_list_path")
    output_dir = intelligence.get("output_dir")
    if isinstance(content_list_path, str) and content_list_path:
        roots.append(Path(content_list_path).expanduser().resolve().parent)
    if isinstance(output_dir, str) and output_dir:
        roots.append(Path(output_dir).expanduser().resolve())

    source_path: Path | None = None
    for root in roots:
        candidate = (root / relative).resolve()
        try:
            candidate.relative_to(root)
        except ValueError:
            continue
        if candidate.is_file() and not _has_link_or_reparse_component(candidate):
            source_path = candidate
            break
    if source_path is None:
        return None

    try:
        source_stat = source_path.stat()
        existing_total = sum(
            int(asset.metadata.get("size_bytes") or 0) for asset in document.assets
        )
        if (
            source_stat.st_size > MAX_SINGLE_ASSET_BYTES
            or existing_total + source_stat.st_size > MAX_DOCUMENT_ASSET_BYTES
        ):
            return None
        with source_path.open("rb") as source:
            asset_sha = hashlib.file_digest(source, "sha256").hexdigest()
        size_bytes = source_stat.st_size
    except OSError:
        return None
    existing = next((asset for asset in document.assets if asset.sha256 == asset_sha), None)
    if existing is not None:
        if span not in existing.source_spans:
            existing.source_spans.append(span)
        return existing.path

    safe_name = _safe_local_asset_name(source_path, asset_sha)
    durable_source = _materialize_adapter_asset(
        document,
        safe_name,
        source_path,
        asset_sha,
        source_kind="PDF/MinerU",
    )
    asset_path = f"assets/{safe_name}"
    document.assets.append(
        CanonicalAsset(
            asset_id=_stable_id("asset", document.document_id, asset_sha),
            path=asset_path,
            media_type=mimetypes.guess_type(source_path.name)[0] or "application/octet-stream",
            sha256=asset_sha,
            source_path=str(durable_source),
            source_spans=[span],
            metadata={
                "source_parser": document.parser_source,
                "mineru_path": raw_path,
                "size_bytes": size_bytes,
                "cache_limit_bytes": ADAPTER_ASSET_CACHE_MAX_BYTES,
                "cache_cleanup_policy": ADAPTER_ASSET_CACHE_CLEANUP_POLICY,
            },
        )
    )
    return asset_path


def _match_pdf_structure(
    entries,
    page_label,
    text: str,
    used: set[str],
    *,
    force: bool = False,
):
    normalized_text = text.strip()
    for entry_page, source_text, structure in entries:
        identifier = next(
            getattr(structure, name)
            for name in ("table_id", "formula_id", "figure_id")
            if hasattr(structure, name)
        )
        if identifier in used or entry_page != page_label:
            continue
        normalized_source = source_text.strip()
        if force or (normalized_source and (
            normalized_source == normalized_text
            or normalized_source in normalized_text
            or normalized_text in normalized_source
        )):
            used.add(identifier)
            return structure
    return None


def _append_unmatched_structures(
    document: CanonicalDocument,
    builder: _DocumentBuilder,
    tables,
    formulas,
    figures,
    used_tables: set[str],
    used_formulas: set[str],
    used_figures: set[str],
) -> None:
    for page_label, text, table in tables:
        if table.table_id not in used_tables and text:
            builder.add_content(
                text,
                table.source_spans[0],
                block_type="table",
                table_id=table.table_id,
            )
    for page_label, text, formula in formulas:
        if formula.formula_id not in used_formulas:
            builder.add_content(
                text,
                formula.source_spans[0],
                block_type="formula",
                formula_id=formula.formula_id,
            )
    for page_label, text, figure in figures:
        if figure.figure_id not in used_figures and text:
            builder.add_content(
                text,
                figure.source_spans[0],
                block_type="figure",
                figure_id=figure.figure_id,
            )


def _is_standalone_pdf_heading(text: str) -> bool:
    stripped = text.strip()
    return bool(re.fullmatch(r"#{1,6}\s+[^\n]+", stripped))


def _extract_explicit_abstract(candidates: list[str]) -> str | None:
    return extract_explicit_abstract(candidates)


def _read_text_layer_for_audit(path: Path, page_count: int) -> tuple[list[str], list[str]]:
    from app.services import parser

    warnings: list[str] = []
    try:
        try:
            page_texts, extracted_count = parser._extract_pdf_text_layer(
                path,
                warnings_out=warnings,
            )
        except TypeError:
            # Preserve compatibility with one-argument integrations and monkeypatches.
            page_texts, extracted_count = parser._extract_pdf_text_layer(path)
    except Exception as exc:  # noqa: BLE001
        page_texts = [""] * page_count
        warnings.append(f"Text layer extraction failed: {exc}")
    else:
        if extracted_count != page_count:
            warnings.append(
                f"Text layer page count {extracted_count} differs from validated count {page_count}."
            )
        if len(page_texts) < page_count:
            page_texts.extend([""] * (page_count - len(page_texts)))
        elif len(page_texts) > page_count:
            page_texts = page_texts[:page_count]
    return page_texts, warnings


def _attach_pdf_audit(
    document: CanonicalDocument,
    *,
    page_count: int,
    page_texts: list[str],
    text_layer_warnings: list[str],
    attempts: list[str],
    primary_parser: str,
    repair_scopes: list[str] | None = None,
) -> None:
    explicit_pages = {
        page
        for page in document.metadata.get("parsed_page_indices", [])
        if isinstance(page, int) and 0 <= page < page_count
    }
    parsed_pages = sorted(
        explicit_pages | _document_structured_page_indices(document)
    )
    blank_text_layer_pages = [
        index
        for index in range(min(page_count, len(page_texts)))
        if not page_texts[index].strip()
    ]
    scopes = list(dict.fromkeys(repair_scopes or document.metadata.get("repair_scopes", [])))
    document.metadata.update(
        {
            "expected_page_count": page_count,
            "parsed_page_indices": parsed_pages,
            "text_layer_pages": list(page_texts),
            "text_layer_blank_page_indices": blank_text_layer_pages,
            "text_layer_warnings": list(text_layer_warnings),
            "primary_parser": primary_parser,
            "parser_attempts": list(attempts),
            "repair_scopes": scopes,
        }
    )
    document.parser_metadata.update(
        {
            "primary_parser": primary_parser,
            "parser_attempts": list(attempts),
            "repair_scopes": scopes,
        }
    )
    for warning in text_layer_warnings:
        if warning not in document.warnings:
            document.warnings.append(warning)


def _finalize_pdf_audit(document: CanonicalDocument) -> CanonicalDocument:
    _finalize_structured_evidence(document)
    page_count = document.metadata.get("expected_page_count")
    page_count = page_count if isinstance(page_count, int) and page_count > 0 else 0
    fallback_pages = set(document.quality.fallback_pages)
    fallback_pages.update(
        page + 1
        for page in _repair_page_indices(
            list(document.metadata.get("repair_scopes", [])),
            page_count,
        )
    )
    if document.metadata.get("primary_parser") != "mineru":
        fallback_pages.update(range(1, page_count + 1))
    quality_dump = document.quality.model_dump(mode="json")
    document.metadata["fallback_pages"] = sorted(fallback_pages)
    document.metadata["quality_status"] = document.quality.status
    document.metadata["quality"] = quality_dump
    document.parser_metadata["fallback_pages"] = sorted(fallback_pages)
    document.parser_metadata["quality_status"] = document.quality.status
    document.parser_metadata["quality"] = quality_dump
    return document


def _repair_page_indices(scopes: list[str], page_count: int) -> set[int]:
    pages: set[int] = set()
    for scope in scopes:
        single = re.fullmatch(r"page:(\d+)", scope)
        if single:
            pages.add(int(single.group(1)) - 1)
            continue
        page_range = re.fullmatch(r"pages:(\d+)-(\d+)", scope)
        if page_range:
            start, end = map(int, page_range.groups())
            pages.update(range(start - 1, end))
    return {page for page in pages if 0 <= page < page_count}


def _repair_scope_page_indices(scope: str) -> set[int]:
    single = re.fullmatch(r"page:(\d+)", scope)
    if single:
        return {int(single.group(1)) - 1}
    page_range = re.fullmatch(r"pages:(\d+)-(\d+)", scope)
    if page_range:
        start, end = map(int, page_range.groups())
        return set(range(start - 1, end)) if end >= start else set()
    return set()


def _structure_on_pages(structure, page_indices: set[int]) -> bool:
    return any(
        span.page_index in page_indices
        for span in structure.source_spans
        if span.page_index is not None
    )


def _block_on_pages(block: CanonicalBlock, page_indices: set[int]) -> bool:
    return any(
        span.page_index in page_indices
        for span in block.source_spans
        if span.page_index is not None
    )


def _document_structured_page_indices(document: CanonicalDocument) -> set[int]:
    structures = [
        *document.blocks,
        *document.tables,
        *document.figures,
        *document.formulas,
    ]
    return {
        span.page_index
        for structure in structures
        for span in structure.source_spans
        if span.page_index is not None
    }


def _targeted_repair_has_complete_coverage(
    repair: CanonicalDocument,
    page_indices: set[int],
) -> bool:
    return bool(page_indices) and page_indices.issubset(
        _document_structured_page_indices(repair)
    )


def _asset_only_on_pages(asset: CanonicalAsset, page_indices: set[int]) -> bool:
    located_pages = {
        span.page_index
        for span in asset.source_spans
        if span.page_index is not None
    }
    return bool(located_pages) and located_pages.issubset(page_indices)


def _reconcile_repair_assets(
    primary_assets: list[CanonicalAsset],
    repair_assets: list[CanonicalAsset],
    figures: list[CanonicalFigure],
    page_indices: set[int],
) -> list[CanonicalAsset]:
    referenced_paths = {
        figure.asset_path for figure in figures if figure.asset_path
    }
    repair_paths = {asset.path for asset in repair_assets}
    candidates = [
        asset.model_copy(deep=True)
        for asset in primary_assets
        if asset.path not in repair_paths
        and not (
            asset.path not in referenced_paths
            and _asset_only_on_pages(asset, page_indices)
        )
    ]
    candidates.extend(
        asset.model_copy(deep=True)
        for asset in repair_assets
        if asset.path in referenced_paths
    )

    by_path: dict[str, CanonicalAsset] = {}
    for asset in candidates:
        existing = by_path.get(asset.path)
        if existing is None:
            by_path[asset.path] = asset
            continue
        for span in asset.source_spans:
            if span not in existing.source_spans:
                existing.source_spans.append(span)
    return list(by_path.values())


def _block_structure_type(block: CanonicalBlock) -> str | None:
    if block.table_id or block.block_type == "table":
        return "table"
    if block.figure_id or block.block_type == "figure":
        return "figure"
    if block.formula_id or block.block_type == "formula":
        return "formula"
    return None


def _issue_replacement_types(issues: list[CanonicalQualityIssue]) -> set[str]:
    replacements: set[str] = set()
    for issue in issues:
        if issue.code.startswith("table_"):
            replacements.add("table")
        elif issue.code.startswith("figure_"):
            replacements.add("figure")
        elif issue.code.startswith("formula_"):
            replacements.add("formula")
    return replacements


def _structures_on_pages(structures: list, page_indices: set[int]) -> list:
    return [
        structure
        for structure in structures
        if _structure_on_pages(structure, page_indices)
    ]


def _structure_bbox(structure) -> tuple[float, float, float, float] | None:
    for span in structure.source_spans:
        box = span.normalized_bbox or span.bbox
        if box is not None:
            return box
    return None


def _table_replacements_match_issues(
    primary: CanonicalDocument,
    replacements: list[CanonicalTable],
    issues: list[CanonicalQualityIssue],
    page_index: int,
) -> bool:
    original_page_tables = _structures_on_pages(primary.tables, {page_index})
    page_issues = [
        issue
        for issue in issues
        if issue.code == "table_invalid"
        and issue.repair_scope
        and page_index in _repair_scope_page_indices(issue.repair_scope)
    ]
    repair_table_ids = {
        str(issue.metadata.get("table_id"))
        for issue in page_issues
        if issue.metadata.get("table_id") is not None
    }
    from app.services.structured_evidence import TableValidator

    bindings = TableValidator().validate_repair_inventory(
        original_page_tables,
        replacements,
        repair_table_ids,
        page_index,
    )
    if bindings is None:
        return False
    for original, replacement, validation, proof in bindings:
        replacement.status = "repaired_by_vision"
        replacement.metadata = {
            **replacement.metadata,
            "repair_original_table_id": original.table_id,
            "repair_proof": proof.model_dump(mode="json"),
            "repair_proof_validated": True,
        }
    return True


def _narrative_blocks_on_pages(
    document: CanonicalDocument,
    page_indices: set[int],
) -> list[CanonicalBlock]:
    return [
        block
        for block in document.blocks
        if _block_on_pages(block, page_indices)
        and _block_structure_type(block) is None
    ]


def _block_fingerprint(block: CanonicalBlock) -> tuple[object, ...]:
    return (
        block.block_id,
        block.block_type,
        block.text,
        tuple(span.model_dump(mode="json") for span in block.source_spans),
    )


def _targeted_repair_satisfies_issues(
    primary: CanonicalDocument,
    repair: CanonicalDocument,
    candidate: CanonicalDocument,
    issues: list[CanonicalQualityIssue],
    page_indices: set[int],
) -> bool:
    from app.services.canonical_quality import CanonicalQualityGate

    unresolved = {
        (issue.code, issue.repair_scope)
        for issue in candidate.quality.issues
        if issue.repairable
    }
    for issue in issues:
        if (issue.code, issue.repair_scope) in unresolved:
            return False
        if issue.code == "abstract_missing" and not (repair.abstract or "").strip():
            return False

    primary_narrative = {
        block.block_id: _block_fingerprint(block)
        for block in _narrative_blocks_on_pages(primary, page_indices)
    }
    candidate_narrative = {
        block.block_id: _block_fingerprint(block)
        for block in _narrative_blocks_on_pages(candidate, page_indices)
    }
    if any(
        candidate_narrative.get(block_id) != fingerprint
        for block_id, fingerprint in primary_narrative.items()
    ):
        return False

    inventories = (
        (primary.tables, repair.tables, candidate.tables, "table"),
        (primary.figures, repair.figures, candidate.figures, "figure"),
        (primary.formulas, repair.formulas, candidate.formulas, "formula"),
    )
    replacement_types = _issue_replacement_types(issues)
    for original, replacements, merged, structure_type in inventories:
        for page_index in sorted(page_indices):
            page = {page_index}
            original_count = len(_structures_on_pages(original, page))
            merged_count = len(_structures_on_pages(merged, page))
            if merged_count < original_count:
                return False
            if structure_type in replacement_types:
                replacement_items = _structures_on_pages(replacements, page)
                if len(replacement_items) != original_count:
                    return False
                if structure_type == "table":
                    replacement_ids = [table.table_id for table in replacement_items]
                    if len(replacement_ids) != len(set(replacement_ids)):
                        return False
                    replacement_fingerprints = [
                        table_identity_fingerprint(table)
                        for table in replacement_items
                    ]
                    if len(replacement_fingerprints) != len(
                        set(replacement_fingerprints)
                    ):
                        return False
                    issue_table_ids = {
                        str(issue.metadata.get("table_id"))
                        for issue in issues
                        if issue.code == "table_invalid"
                        and issue.metadata.get("table_id") is not None
                        and issue.repair_scope
                        and page_index in _repair_scope_page_indices(issue.repair_scope)
                    }
                    if issue_table_ids and len(replacement_items) < len(issue_table_ids):
                        return False
                    if not _table_replacements_match_issues(
                        primary,
                        replacement_items,
                        issues,
                        page_index,
                    ):
                        return False
                if structure_type == "table" and any(
                    CanonicalQualityGate._invalid_table_reasons(table)
                    for table in replacement_items
                ):
                    return False
    return True


def _merge_pdf_page_repairs(
    primary: CanonicalDocument,
    repair: CanonicalDocument,
    page_indices: set[int],
    issues: list[CanonicalQualityIssue] | None = None,
) -> CanonicalDocument:
    replacement_types = (
        {"table", "figure", "formula"}
        if issues is None
        else _issue_replacement_types(issues)
    )
    # Targeted repairs replace only the explicitly requested structures.  In
    # particular, abstract recovery updates the document-level field and does
    # not replace ordinary narrative blocks on the target pages.
    replace_page_text = issues is None

    def replace_block(block: CanonicalBlock) -> bool:
        if not _block_on_pages(block, page_indices):
            return False
        structure_type = _block_structure_type(block)
        if structure_type is not None:
            return structure_type in replacement_types
        return replace_page_text

    original_blocks = list(primary.blocks)
    removed_blocks = [block for block in original_blocks if replace_block(block)]
    retained_blocks = [block for block in original_blocks if not replace_block(block)]
    incoming_blocks = [
        block.model_copy(deep=True)
        for block in repair.blocks
        if (
            (_block_structure_type(block) in replacement_types)
            or (_block_structure_type(block) is None and replace_page_text)
        )
    ]

    def block_page(block: CanonicalBlock) -> int | None:
        return next(
            (
                span.page_index
                for span in block.source_spans
                if span.page_index is not None
            ),
            None,
        )

    def block_kind(block: CanonicalBlock) -> str:
        return _block_structure_type(block) or block.block_type

    def block_position(block: CanonicalBlock) -> tuple[float, ...] | None:
        span = next(iter(block.source_spans), None)
        if span is None:
            return None
        box = span.normalized_bbox or span.bbox
        if box is not None:
            return (
                box[1],
                box[0],
                float(span.line_start or 0),
                float(span.char_start or 0),
            )
        if span.line_start is not None or span.char_start is not None:
            return (float(span.line_start or 0), float(span.char_start or 0))
        return None

    def group(block: CanonicalBlock) -> tuple[int | None, str]:
        return block_page(block), block_kind(block)

    anchors: dict[str, int] = {
        block.block_id: block.reading_order for block in retained_blocks
    }
    block_id_map: dict[str, str] = {}
    for key in {
        group(block)
        for block in removed_blocks
    }:
        old_group = sorted(
            (block for block in removed_blocks if group(block) == key),
            key=lambda block: block.reading_order,
        )
        new_group = sorted(
            (block for block in incoming_blocks if group(block) == key),
            key=lambda block: block.reading_order,
        )
        for old_block, new_block in zip(old_group, new_group):
            block_id_map[old_block.block_id] = new_block.block_id
            anchors[new_block.block_id] = old_block.reading_order
    for block in incoming_blocks:
        anchors.setdefault(block.block_id, block.reading_order)

    retained_ids = {block.block_id for block in retained_blocks}
    occupied_ids = set(retained_ids)
    occupied_ids.update(block.block_id for block in incoming_blocks)
    seen_incoming: set[str] = set()
    for block in incoming_blocks:
        original_id = block.block_id
        if original_id in seen_incoming or original_id in retained_ids:
            block.block_id = _stable_id(
                "block", primary.document_id, "repair", len(seen_incoming), original_id
            )
            while block.block_id in occupied_ids:
                block.block_id = _stable_id(
                    "block", primary.document_id, "repair", len(occupied_ids), block.block_id
                )
            for old_id, replacement_id in list(block_id_map.items()):
                if replacement_id == original_id:
                    block_id_map[old_id] = block.block_id
            anchors[block.block_id] = anchors.pop(original_id, block.reading_order)
        seen_incoming.add(block.block_id)
        occupied_ids.add(block.block_id)

    primary.blocks = retained_blocks + incoming_blocks
    if "table" in replacement_types:
        primary.tables = [
            table for table in primary.tables if not _structure_on_pages(table, page_indices)
        ] + [table.model_copy(deep=True) for table in repair.tables]
    if "figure" in replacement_types:
        primary.figures = [
            figure for figure in primary.figures if not _structure_on_pages(figure, page_indices)
        ] + [figure.model_copy(deep=True) for figure in repair.figures]
        primary.assets = _reconcile_repair_assets(
            primary.assets,
            repair.assets,
            primary.figures,
            page_indices,
        )
    if "formula" in replacement_types:
        primary.formulas = [
            formula for formula in primary.formulas if not _structure_on_pages(formula, page_indices)
        ] + [formula.model_copy(deep=True) for formula in repair.formulas]
    if repair.abstract and repair.abstract.strip():
        primary.abstract = repair.abstract

    primary.blocks.sort(
        key=lambda block: (
            block_page(block) if block_page(block) is not None else 10**9,
            0 if block_position(block) is not None else 1,
            *(block_position(block) or ()),
            anchors.get(block.block_id, block.reading_order),
            block.block_id,
        )
    )
    final_block_ids = {block.block_id for block in primary.blocks}
    for figure in primary.figures:
        figure.nearby_block_ids = [
            block_id_map.get(block_id, block_id)
            for block_id in figure.nearby_block_ids
            if block_id_map.get(block_id, block_id) in final_block_ids
        ]
    for formula in primary.formulas:
        formula.nearby_block_ids = [
            block_id_map.get(block_id, block_id)
            for block_id in formula.nearby_block_ids
            if block_id_map.get(block_id, block_id) in final_block_ids
        ]

    def remap_outline(nodes: list[SectionNode]) -> list[SectionNode]:
        remapped: list[SectionNode] = []
        for node in nodes:
            mapped_id = block_id_map.get(node.block_id, node.block_id) if node.block_id else None
            if node.block_id and mapped_id not in final_block_ids:
                continue
            remapped.append(
                node.model_copy(
                    update={
                        "block_id": mapped_id,
                        "children": remap_outline(node.children),
                    },
                    deep=True,
                )
            )
        return remapped

    primary.outline = remap_outline(primary.outline)
    for issue in primary.quality.issues:
        issue.block_ids = [
            block_id_map.get(block_id, block_id)
            for block_id in issue.block_ids
            if block_id_map.get(block_id, block_id) in final_block_ids
        ]
    for reading_order, block in enumerate(primary.blocks):
        block.reading_order = reading_order
    return primary


class PDFCanonicalAdapter:
    parser_source = "mineru"

    def parse(self, path: Path) -> CanonicalDocument:
        path = _validate_path(path)
        from app.services import parser
        from app.services.canonical_quality import CanonicalQualityGate

        page_count = parser._validate_pdf_basic(path)
        attempts: list[str] = []

        mineru_document: CanonicalDocument | None = None
        if parser.settings.mineru_enabled:
            try:
                mineru_document = run_mineru(path, page_count)
                attempts.append(
                    "mineru:success" if mineru_document is not None else "mineru:unavailable"
                )
            except Exception as exc:  # noqa: BLE001
                attempts.append(f"mineru:failed:{type(exc).__name__}:{exc}")
        else:
            attempts.append("mineru:disabled")

        page_texts, text_layer_warnings = _read_text_layer_for_audit(path, page_count)
        gate = CanonicalQualityGate()
        if mineru_document is not None:
            _attach_pdf_audit(
                mineru_document,
                page_count=page_count,
                page_texts=page_texts,
                text_layer_warnings=text_layer_warnings,
                attempts=attempts,
                primary_parser="mineru",
            )
            report = gate.evaluate(mineru_document)
            _finalize_structured_evidence(mineru_document)
            report = mineru_document.quality
            if report.accepted:
                return _finalize_pdf_audit(mineru_document)

            fatal = any(issue.severity == "fatal" for issue in report.issues)
            repair_scopes = [
                issue.repair_scope
                for issue in report.issues
                if issue.repairable and issue.repair_scope
            ]
            repair_issues = [issue for issue in report.issues if issue.repairable]
            targeted_pages = _repair_page_indices(repair_scopes, page_count)
            if not fatal and repair_scopes and targeted_pages and parser.settings.document_intelligence_enabled:
                try:
                    repair = run_document_intelligence(
                        path,
                        page_count,
                        page_texts,
                        page_indices=targeted_pages,
                    )
                except Exception as exc:  # noqa: BLE001
                    repair = None
                    attempts.append(
                        f"document_intelligence:targeted:failed:{type(exc).__name__}:{exc}"
                    )
                else:
                    attempts.append(
                        "document_intelligence:targeted:success"
                        if repair is not None
                        else "document_intelligence:targeted:unavailable"
                    )
                if repair is not None:
                    if not _targeted_repair_has_complete_coverage(
                        repair,
                        targeted_pages,
                    ):
                        attempts.append("document_intelligence:targeted:partial")
                        _attach_pdf_audit(
                            mineru_document,
                            page_count=page_count,
                            page_texts=page_texts,
                            text_layer_warnings=text_layer_warnings,
                            attempts=attempts,
                            primary_parser="mineru",
                        )
                        return _finalize_pdf_audit(mineru_document)
                    if "table" in _issue_replacement_types(repair_issues) and any(
                        not _table_replacements_match_issues(
                            mineru_document,
                            _structures_on_pages(repair.tables, {page_index}),
                            repair_issues,
                            page_index,
                        )
                        for page_index in sorted(targeted_pages)
                    ):
                        attempts.append("document_intelligence:targeted:incomplete")
                        _attach_pdf_audit(
                            mineru_document,
                            page_count=page_count,
                            page_texts=page_texts,
                            text_layer_warnings=text_layer_warnings,
                            attempts=attempts,
                            primary_parser="mineru",
                        )
                        return _finalize_pdf_audit(mineru_document)
                    candidate = _merge_pdf_page_repairs(
                        mineru_document.model_copy(deep=True),
                        repair,
                        targeted_pages,
                        issues=repair_issues,
                    )
                    _attach_pdf_audit(
                        candidate,
                        page_count=page_count,
                        page_texts=page_texts,
                        text_layer_warnings=text_layer_warnings,
                        attempts=attempts,
                        primary_parser="mineru",
                        repair_scopes=repair_scopes,
                    )
                    repaired_report = gate.evaluate(candidate)
                    if (
                        not any(
                            issue.severity == "fatal"
                            for issue in repaired_report.issues
                        )
                        and _targeted_repair_satisfies_issues(
                            mineru_document,
                            repair,
                            candidate,
                            repair_issues,
                            targeted_pages,
                        )
                    ):
                        return _finalize_pdf_audit(candidate)
                    attempts.append("document_intelligence:targeted:incomplete")
                    _attach_pdf_audit(
                        mineru_document,
                        page_count=page_count,
                        page_texts=page_texts,
                        text_layer_warnings=text_layer_warnings,
                        attempts=attempts,
                        primary_parser="mineru",
                    )
                    return _finalize_pdf_audit(mineru_document)
                else:
                    _attach_pdf_audit(
                        mineru_document,
                        page_count=page_count,
                        page_texts=page_texts,
                        text_layer_warnings=text_layer_warnings,
                        attempts=attempts,
                        primary_parser="mineru",
                    )
                    return _finalize_pdf_audit(mineru_document)
            elif not fatal:
                if repair_scopes and not parser.settings.document_intelligence_enabled:
                    attempts.append("document_intelligence:targeted:disabled")
                return _finalize_pdf_audit(mineru_document)

        if parser.settings.document_intelligence_enabled:
            try:
                full_di = run_document_intelligence(path, page_count, page_texts)
            except Exception as exc:  # noqa: BLE001
                full_di = None
                attempts.append(
                    f"document_intelligence:full:failed:{type(exc).__name__}:{exc}"
                )
            else:
                attempts.append(
                    "document_intelligence:full:success"
                    if full_di is not None
                    else "document_intelligence:full:unavailable"
                )
            if full_di is not None:
                _attach_pdf_audit(
                    full_di,
                    page_count=page_count,
                    page_texts=page_texts,
                    text_layer_warnings=text_layer_warnings,
                    attempts=attempts,
                    primary_parser="document_intelligence",
                )
                report = gate.evaluate(full_di)
                if not any(issue.severity == "fatal" for issue in report.issues):
                    return _finalize_pdf_audit(full_di)
        else:
            attempts.append("document_intelligence:full:disabled")

        fallback = run_text_layer_fallback(
            path,
            page_count,
            page_texts,
            text_layer_warnings=text_layer_warnings,
        )
        attempts.append("pypdf_text_layer:success")
        _attach_pdf_audit(
            fallback,
            page_count=page_count,
            page_texts=page_texts,
            text_layer_warnings=text_layer_warnings,
            attempts=attempts,
            primary_parser="pypdf_text_layer",
        )
        gate.evaluate(fallback)
        return _finalize_pdf_audit(fallback)


def parse_canonical_document(path: Path) -> CanonicalDocument:
    path = _validate_path(Path(path))
    adapter: CanonicalAdapter = {
        ".pdf": PDFCanonicalAdapter(),
        ".docx": DocxCanonicalAdapter(),
        ".html": HtmlCanonicalAdapter(),
        ".htm": HtmlCanonicalAdapter(),
        ".md": MarkdownCanonicalAdapter(),
        ".markdown": MarkdownCanonicalAdapter(),
        ".txt": TextCanonicalAdapter(),
    }.get(path.suffix.lower(), TextCanonicalAdapter())
    document = adapter.parse(path)
    if isinstance(document, CanonicalDocument):
        _finalize_structured_evidence(document)
        _link_nearby_structured_source_blocks(document)
    return document


def _finalize_structured_evidence(document: CanonicalDocument) -> None:
    """Merge and validate tables at the canonical adapter boundary."""

    from app.services.canonical_quality import CanonicalQualityGate
    from app.services.structured_evidence import StructuredEvidenceBuilder, TableValidator

    # Tokenization is not used by continuation merging. Injecting this counter
    # keeps adapter finalization independent from model-cache availability.
    builder = StructuredEvidenceBuilder(token_counter=lambda _text: 0)
    document.tables = builder.merge_cross_page_tables(document.tables)
    aliases = dict(builder.table_aliases)
    if aliases:
        document.metadata["table_aliases"] = aliases

        def remap_metadata(value):
            if isinstance(value, list):
                return [remap_metadata(item) for item in value]
            if not isinstance(value, dict):
                return value
            remapped = {}
            for key, item in value.items():
                if (
                    key
                    in {
                        "table_id",
                        "parent_table_id",
                        "continuation_of",
                        "original_table_id",
                        "retrieval_covered_by_table_id",
                    }
                    and isinstance(item, str)
                ):
                    remapped[key] = aliases.get(item, item)
                else:
                    remapped[key] = remap_metadata(item)
            return remapped

        for block in document.blocks:
            if block.table_id is not None:
                block.table_id = aliases.get(block.table_id, block.table_id)
            block.metadata = remap_metadata(block.metadata)
            for span in block.source_spans:
                if span.table_id is not None:
                    span.table_id = aliases.get(span.table_id, span.table_id)
                span.metadata = remap_metadata(span.metadata)
        for issue in document.quality.issues:
            issue.metadata = remap_metadata(issue.metadata)
    validator = TableValidator()
    requests: list[dict[str, object]] = []
    validated_tables: list[CanonicalTable] = []
    failed_results = []
    for table in document.tables:
        result = validator.validate(table)
        if result.accepted:
            result.table.metadata.pop("structured_validation_reasons", None)
        else:
            result.table.metadata["structured_validation_reasons"] = list(
                result.reasons
            )
        validated_tables.append(result.table)
        if not result.accepted:
            failed_results.append(result)
            if result.repair_request is not None:
                requests.append(result.repair_request.model_dump(mode="json"))
    document.tables = validated_tables
    document.metadata["table_repair_requests"] = requests
    document.metadata["table_activation_allowed"] = not failed_results

    CanonicalQualityGate().evaluate(document)
    failed_by_id = {result.table.table_id: result for result in failed_results}
    for issue in document.quality.issues:
        if issue.code != "table_invalid":
            continue
        result = failed_by_id.get(str(issue.metadata.get("table_id")))
        if result is None:
            continue
        repair_request = result.repair_request
        issue.metadata["reasons"] = list(result.reasons)
        issue.metadata["repair_request"] = (
            repair_request.model_dump(mode="json")
            if repair_request is not None
            else None
        )


def _link_nearby_structured_source_blocks(document: CanonicalDocument) -> None:
    """Attach nearby source prose without manufacturing descriptions.

    The links are format-neutral: every adapter already emits canonical blocks,
    so Figure/Formula retrieval can use the same provenance rule for PDF, DOCX,
    HTML, Markdown, and plain text.
    """

    ordered = sorted(document.blocks, key=lambda block: (block.reading_order, block.block_id))
    positions = {block.block_id: index for index, block in enumerate(ordered)}

    def block_pages(block: CanonicalBlock) -> set[int]:
        return {
            span.page_index
            for span in block.source_spans
            if span.page_index is not None
        }

    def nearby_ids(
        owner_id: str,
        owner_pages: set[int],
        existing: list[str],
        attribute: str,
    ) -> list[str]:
        anchors = [
            positions[block.block_id]
            for block in ordered
            if getattr(block, attribute) == owner_id
        ]
        if not anchors:
            return existing
        candidates: list[tuple[int, int, str]] = []
        for block in ordered:
            if (
                block.block_type not in {"narrative", "appendix"}
                or not block.text.strip()
                or block_is_generated(block)
            ):
                continue
            pages = block_pages(block)
            if owner_pages and pages and owner_pages.isdisjoint(pages):
                continue
            position = positions[block.block_id]
            distance = min(abs(position - anchor) for anchor in anchors)
            if distance <= 2:
                candidates.append((distance, position, block.block_id))
        discovered = [item[2] for item in sorted(candidates)[:2]]
        return list(dict.fromkeys([*existing, *discovered]))

    for figure in document.figures:
        pages = {
            span.page_index
            for span in figure.source_spans
            if span.page_index is not None
        }
        figure.nearby_block_ids = nearby_ids(
            figure.figure_id,
            pages,
            figure.nearby_block_ids,
            "figure_id",
        )
    for formula in document.formulas:
        pages = {
            span.page_index
            for span in formula.source_spans
            if span.page_index is not None
        }
        formula.nearby_block_ids = nearby_ids(
            formula.formula_id,
            pages,
            formula.nearby_block_ids,
            "formula_id",
        )
